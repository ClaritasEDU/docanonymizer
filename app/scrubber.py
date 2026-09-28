"""Apply replacements + deep scrub for DOCX/XLSX (PRD 5.9).

Two-stage strategy:

1. **Surface replacement**: walk the document model with the format-native
   library and replace text where it lives (paragraph runs, table cells,
   workbook cells, slide shapes). This preserves formatting.

2. **Deep scrub**: unzip the resulting DOCX/XLSX, do a raw-XML string
   substitution pass across every part, strip tracked changes /
   comments / author metadata. Repack as a fresh archive. The output is
   never a modified copy of the original binary.

Key scrub layers per format:

  DOCX
    - Paragraph runs, table cells, headers, footers (python-docx pass)
    - Walk every part in the .docx zip; raw map substitution
    - Strip <w:ins>/<w:del> revision elements (keep final text)
    - Clear word/comments.xml
    - Zero docProps/core.xml: <dc:creator>, <cp:lastModifiedBy>
    - Strip alt text on images (wp:docPr/@descr, pic:cNvPr/@descr)

  XLSX
    - Cell values via openpyxl (preserves formatting/formulas)
    - Walk every part in the .xlsx zip; raw map substitution
    - Clear xl/comments*.xml
    - Zero docProps/core.xml: <dc:creator>, <cp:lastModifiedBy>
    - Flag formula strings that contain a matched PII value (do not silently scrub)

Every writer takes either a replacement map (original -> placeholder) or a
Replacer (replacer.py). Matching is single-pass and never rewrites inside an
existing placeholder. Unanonymize reuses these same writers with a
TokenRestorer and `deep_clean=False` - restoring a file must not strip the
comments or metadata of the document the AI handed back.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, Union

from . import ids
from .extractors import ExtractResult
from .logging_setup import get_logger
from .replacer import Replacer, apply, as_replacer, is_short_number, splice, xml_escape

MapOrReplacer = Union[dict, Replacer]

log = get_logger("scrubber")


@dataclass
class ScrubResult:
    output_path: Path
    layers_cleaned: list[str] = field(default_factory=list)
    formula_warnings: list[str] = field(default_factory=list)
    bytes_written: int = 0
    # XLSX: anonymized sheet title -> original title, stored in the key file
    # so restore puts the exact title back (titles are length-limited).
    sheet_titles: dict[str, str] = field(default_factory=dict)


def apply_replacements_text(text: str, replacement_map: MapOrReplacer) -> str:
    """Single-pass substitution (leftmost-longest, placeholders protected).

    Passing a dict compiles a matcher on every call - hot loops should call
    `as_replacer()` once and pass the Replacer in.
    """
    return apply(text, as_replacer(replacement_map))


# ---------------------------------------------------------------------------
# Run-aware XML replacement (CODE_REVIEW C3)
#
# Word routinely splits a single string like "Jane Smith" across multiple
# <w:r>/<w:t> runs after edits, so neither a per-run pass nor a raw string
# substitution can see it. These helpers join the text runs of each paragraph
# (or shared-string item), find matches in the joined text, and write the
# replacement back into the run where the match starts - removing the matched
# characters from the following runs. Formatting outside the match survives.
# The same pass restores split identifiers in unanonymize.
# ---------------------------------------------------------------------------

_WP_BLOCK_RE = re.compile(r"<w:p[ >].*?</w:p>", re.DOTALL)
_WT_RE = re.compile(r"<w:t((?:\s[^>]*)?)/>|<w:t((?:\s[^>]*)?)>(.*?)</w:t>", re.DOTALL)
_SI_BLOCK_RE = re.compile(r"<si>.*?</si>", re.DOTALL)
_IS_BLOCK_RE = re.compile(r"<is>.*?</is>", re.DOTALL)
_T_RE = re.compile(r"<t((?:\s[^>]*)?)/>|<t((?:\s[^>]*)?)>(.*?)</t>", re.DOTALL)

_DOCX_RUN_PARTS = re.compile(r"word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml")
_XLSX_SHEET_PARTS = re.compile(r"xl/worksheets/sheet\d*\.xml")

_ENTITY_RE = re.compile(r"&(amp|lt|gt|quot|apos|#x?[0-9a-fA-F]+);")


def _xml_unescape(s: str) -> str:
    def repl(m: re.Match) -> str:
        e = m.group(1)
        table = {"amp": "&", "lt": "<", "gt": ">", "quot": '"', "apos": "'"}
        if e in table:
            return table[e]
        try:
            code = int(e[2:], 16) if e[1] in "xX" else int(e[1:])
            return chr(code)
        except (ValueError, OverflowError):
            return m.group(0)
    return _ENTITY_RE.sub(repl, s)


_xml_escape = xml_escape

_CHAR_REF_RE = re.compile(r"&#(x[0-9A-Fa-f]+|[0-9]+);")


def _decode_char_refs(xml: str) -> str:
    """Turn numeric character references for non-ASCII characters into the
    characters themselves: openpyxl writes "José" as "Jos&#233;". Without
    this, no text pass can match an accented or non-Latin name in formula
    text, metadata, or any other raw XML - and it would ship. ASCII
    references (&#10;, &#38;...) are left alone: they can carry meaning."""
    def repl(m: re.Match) -> str:
        ref = m.group(1)
        try:
            cp = int(ref[1:], 16) if ref[0] in "xX" else int(ref)
        except ValueError:
            return m.group(0)
        if cp < 0x80 or cp > 0x10FFFF or 0xD800 <= cp <= 0xDFFF:
            return m.group(0)
        return chr(cp)
    return _CHAR_REF_RE.sub(repl, xml)


def _replace_in_blocks(xml: str, block_re: re.Pattern, t_re: re.Pattern,
                       t_name: str, replacer: Replacer) -> str:
    """Apply the replacer inside each text-run block, spanning split runs."""

    def process_block(bm: re.Match) -> str:
        block = bm.group(0)
        nodes = []  # (start, end, attrs, text)
        for m in t_re.finditer(block):
            attrs = m.group(1) if m.group(1) is not None else (m.group(2) or "")
            inner = m.group(3) if m.group(3) is not None else ""
            nodes.append((m.start(), m.end(), attrs, _xml_unescape(inner)))
        if not nodes:
            return block
        joined = "".join(n[3] for n in nodes)
        if not joined:
            return block

        ranges = replacer.find(joined)
        if not ranges:
            return block

        # Map each joined-text character back to its owning node.
        owner: list[int] = []
        for i, n in enumerate(nodes):
            owner.extend([i] * len(n[3]))

        new_texts = [""] * len(nodes)
        i = 0
        r = 0
        while i < len(joined):
            if r < len(ranges) and ranges[r][0] == i:
                new_texts[owner[i]] += ranges[r][2]
                i = ranges[r][1]
                r += 1
            else:
                new_texts[owner[i]] += joined[i]
                i += 1

        # Rebuild the block back-to-front so offsets stay valid.
        rebuilt = block
        for k in range(len(nodes) - 1, -1, -1):
            start, end, attrs, _ = nodes[k]
            if "xml:space" not in attrs:
                attrs += ' xml:space="preserve"'
            node_xml = f"<{t_name}{attrs}>{_xml_escape(new_texts[k])}</{t_name}>"
            rebuilt = rebuilt[:start] + node_xml + rebuilt[end:]
        return rebuilt

    return block_re.sub(process_block, xml)


def _structured_replace_docx(name: str, xml: str, replacer: Replacer) -> str:
    if _DOCX_RUN_PARTS.fullmatch(name):
        return _replace_in_blocks(xml, _WP_BLOCK_RE, _WT_RE, "w:t", replacer)
    return xml


def _structured_replace_xlsx(name: str, xml: str, replacer: Replacer) -> str:
    if name == "xl/sharedStrings.xml":
        return _replace_in_blocks(xml, _SI_BLOCK_RE, _T_RE, "t", replacer)
    if _XLSX_SHEET_PARTS.fullmatch(name):
        return _replace_in_blocks(xml, _IS_BLOCK_RE, _T_RE, "t", replacer)
    return xml


# ---------------------------------------------------------------------------
# Format-specific writers
# ---------------------------------------------------------------------------

def scrub_text(extract: ExtractResult, replacement_map: MapOrReplacer, out_path: Path) -> ScrubResult:
    new_text = apply(extract.text, as_replacer(replacement_map))
    out_path.write_text(new_text, encoding="utf-8")
    return ScrubResult(
        output_path=out_path,
        layers_cleaned=["text"],
        bytes_written=out_path.stat().st_size,
    )


def _override_for(cell_overrides: Optional[dict], key: tuple, current: str):
    """The replacement for this cell, or None.

    `cell_overrides` maps (sheet, row, col) -> (expected, new): the cell is
    rewritten to `new` only if it still holds exactly `expected` (None = any
    value). Anonymize uses it to put community identifiers into name cells;
    restore uses it to put the exact original text back - and only where the
    token is still there untouched."""
    if not cell_overrides:
        return None
    hit = cell_overrides.get(key)
    if hit is None:
        return None
    expected, new = hit
    if expected is not None and current != expected:
        log.warning("cell override skipped at sheet=%d row=%d col=%d: content changed", *key)
        return None
    return new


def scrub_csv(extract: ExtractResult, replacement_map: MapOrReplacer, out_path: Path,
              cell_overrides: Optional[dict] = None,
              append_columns: Optional[list] = None,
              remove_columns: Optional[list] = None) -> ScrubResult:
    replacer = as_replacer(replacement_map)
    rows = extract.payload.get("rows") or []
    # Judge an appended column on the file as it came in - after the token
    # pass its family ids have already become household names.
    removable = []
    for spec in remove_columns or []:
        if _appended_column_intact(rows, spec):
            removable.append(spec)
        else:
            log.warning("appended %s column not removed: it was edited", spec.get("header"))
    new_rows: list[list[str]] = []
    for r, row in enumerate(rows):
        out = []
        for c, cell in enumerate(row):
            forced = _override_for(cell_overrides, (0, r, c), cell)
            out.append(forced if forced is not None else apply(cell, replacer))
        new_rows.append(out)
    for spec in append_columns or []:
        col = spec["col"]
        for r, row in enumerate(new_rows):
            if r == spec["header_row"]:
                value = spec["header"]
            else:
                value = spec["values"].get(r, "")
            if len(row) < col:
                row.extend([""] * (col - len(row)))
            row.insert(col, value)
    for spec in sorted(removable, key=lambda s: -s["col"]):
        for row in new_rows:
            if len(row) > spec["col"]:
                del row[spec["col"]]
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(new_rows)
    return ScrubResult(
        output_path=out_path,
        layers_cleaned=["cells"],
        bytes_written=out_path.stat().st_size,
    )


def scrub_docx(working_path: Path, replacement_map: MapOrReplacer, out_path: Path,
               deep_clean: bool = True) -> ScrubResult:
    """DOCX scrub: surface replace via python-docx, then ZIP-level deep scrub.

    `deep_clean=False` (unanonymize) still runs every replacement pass but
    leaves comments, tracked changes, alt text, and metadata alone.
    """
    from docx import Document

    replacer = as_replacer(replacement_map)

    # Stage 1 - format-aware surface replacement on a temp copy.
    staged = out_path.with_suffix(".staged.docx")
    shutil.copyfile(working_path, staged)
    doc = Document(str(staged))

    def _replace_runs(runs):
        for run in runs:
            if run.text:
                new = apply(run.text, replacer)
                if new != run.text:
                    run.text = new

    for section in doc.sections:
        for box in (section.header, section.footer):
            for p in box.paragraphs:
                _replace_runs(p.runs)
    for p in doc.paragraphs:
        _replace_runs(p.runs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_runs(p.runs)
    doc.save(str(staged))

    # Stage 2 - ZIP-level deep scrub (run-aware pass handles split runs).
    layers = _deep_scrub_zip(
        staged,
        out_path,
        replacer=replacer,
        structured_replace=_structured_replace_docx,
        per_part_handlers={
            "word/document.xml":   _strip_revision_marks,
            "word/header*.xml":    _strip_revision_marks,
            "word/footer*.xml":    _strip_revision_marks,
            "word/comments.xml":   _empty_comments_xml,
            "word/commentsExtended.xml": _empty_comments_ex_xml,
            "docProps/core.xml":   _zero_core_authors,
            "docProps/app.xml":    _zero_app_company,
        } if deep_clean else {},
    )
    staged.unlink(missing_ok=True)

    return ScrubResult(
        output_path=out_path,
        layers_cleaned=["docx_runs"] + layers,
        bytes_written=out_path.stat().st_size,
    )


def _appended_column_intact(rows: list, spec: dict) -> bool:
    """Only remove a column we appended if it still is exactly that column:
    the header in place and every other cell empty or a family token."""
    from .community import COMMUNITY_TOKEN_RE
    col = spec["col"]
    for r, row in enumerate(rows):
        v = row[col] if col < len(row) else ""
        v = "" if v is None else str(v)
        if r == spec["header_row"]:
            if v != spec["header"]:
                return False
        elif v and not COMMUNITY_TOKEN_RE.fullmatch(v):
            return False
    return True


def scrub_xlsx(working_path: Path, replacement_map: MapOrReplacer, out_path: Path,
               deep_clean: bool = True,
               sheet_restore: Optional[dict[str, str]] = None,
               cell_overrides: Optional[dict] = None,
               append_columns: Optional[list] = None,
               remove_columns: Optional[list] = None) -> ScrubResult:
    """XLSX scrub: surface replace via openpyxl, ZIP-level deep scrub, formula warnings.

    `deep_clean=False` (unanonymize) skips the comment/metadata handlers.
    `sheet_restore` (unanonymize) maps anonymized sheet titles back to the
    originals recorded in the key file.
    `cell_overrides` / `append_columns` / `remove_columns`: community
    identifiers (see community.py and _override_for). Coordinates are
    0-based (sheet index, row, column) as extracted.
    """
    from openpyxl import load_workbook

    replacer = as_replacer(replacement_map)

    staged = out_path.with_suffix(".staged.xlsx")
    shutil.copyfile(working_path, staged)
    wb = load_workbook(str(staged))

    # Judge an appended column on the file as it came in - after the token
    # pass its family ids have already become household names.
    removable = []
    for spec in remove_columns or []:
        ws0 = wb.worksheets[spec["sheet"]] if spec["sheet"] < len(wb.worksheets) else None
        grid = [list(r) for r in ws0.iter_rows(values_only=True)] if ws0 is not None else []
        if ws0 is not None and ws0.max_column == spec["col"] + 1 and _appended_column_intact(grid, spec):
            removable.append(spec)
        else:
            log.warning("appended %s column not removed: it was edited or moved", spec.get("header"))

    formula_warnings: list[str] = []
    for si, ws in enumerate(wb.worksheets):
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if cell_overrides:
                    current = "" if v is None else str(v)
                    forced = _override_for(cell_overrides, (si, cell.row - 1, cell.column - 1), current)
                    if forced is not None:
                        cell.value = forced
                        continue
                if v is None:
                    continue
                if isinstance(v, str):
                    if v.startswith("="):
                        # Formula - flag if a PII string is embedded; don't silently rewrite.
                        if deep_clean and replacer.find(v):
                            formula_warnings.append(
                                f"sheet={ws.title!r} cell={cell.coordinate} formula references PII"
                            )
                    else:
                        new = apply(v, replacer)
                        if new != v:
                            cell.value = new
                elif deep_clean and _is_numeric_or_date(v):
                    # A phone typed without dashes, a ZIP, or a date of birth is
                    # stored as a number/date, not text. Match against the same
                    # str() form the extractor showed the model; a hit turns the
                    # cell into text. Skipping these used to leave the value in
                    # place for the raw XML pass to jam a placeholder into a
                    # numeric cell - a workbook Excel has to "repair".
                    sv = str(v)
                    new = apply(sv, replacer)
                    if new != sv:
                        cell.value = new

    # Sheet names can carry PII ("Smith Family") but can't contain [ or ] and
    # max out at 31 characters, so they get the bracket-free token form (or a
    # unique SHEET_<id> if that won't fit). The key file records every rename,
    # so restore is exact. References to the old name in formulas, defined
    # names, and pivot sources are rewritten to match.
    for spec in append_columns or []:
        ws = wb.worksheets[spec["sheet"]]
        ws.cell(row=spec["header_row"] + 1, column=spec["col"] + 1).value = spec["header"]
        for r, token in spec["values"].items():
            ws.cell(row=int(r) + 1, column=spec["col"] + 1).value = token
    for spec in removable:
        wb.worksheets[spec["sheet"]].delete_cols(spec["col"] + 1)

    if sheet_restore:
        renames = {cur: orig for cur, orig in sheet_restore.items() if cur in wb.sheetnames}
    elif deep_clean:
        renames = _anonymized_sheet_titles(wb, replacer)
    else:
        renames = {}
    _apply_sheet_renames(wb, renames)
    wb.save(str(staged))

    def structured(name: str, xml: str, r: Replacer) -> str:
        if renames:
            xml = _rewrite_sheet_refs(xml, renames)
        return _structured_replace_xlsx(name, xml, r)

    layers = _deep_scrub_zip(
        staged,
        out_path,
        replacer=replacer,
        structured_replace=structured,
        raw_protect=_protect_cell_values,
        per_part_handlers={
            re.compile(r"xl/comments\d*\.xml"): _empty_xlsx_comments_xml,
            "docProps/core.xml": _zero_core_authors,
            "docProps/app.xml":  _zero_app_company,
        } if deep_clean else {},
    )
    staged.unlink(missing_ok=True)

    if formula_warnings:
        log.warning("xlsx formula warnings: %d", len(formula_warnings))

    if renames:
        layers.append("sheet_names")
    return ScrubResult(
        output_path=out_path,
        layers_cleaned=["xlsx_cells"] + layers,
        formula_warnings=formula_warnings,
        bytes_written=out_path.stat().st_size,
        sheet_titles={new: old for old, new in renames.items()} if deep_clean else {},
    )


def _is_numeric_or_date(v) -> bool:
    if isinstance(v, bool):
        return False
    return isinstance(v, (int, float, dt.date, dt.time, dt.timedelta))


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\']")
_BRACKETED_ID = re.compile(r"\[([A-Z]+_[0-9A-F]{12})\]")


def _anonymized_sheet_titles(wb, replacer: Replacer) -> dict[str, str]:
    """old title -> new title for every sheet whose title contains PII."""
    renames: dict[str, str] = {}
    taken = {ws.title.casefold() for ws in wb.worksheets}   # Excel: case-insensitive
    reserved: Optional[set[str]] = None
    for ws in wb.worksheets:
        old = ws.title
        new = apply(old, replacer)
        if new == old:
            continue
        new = _BAD_SHEET_CHARS.sub("_", _BRACKETED_ID.sub(r"\1", new)).strip()
        if not new or len(new) > 31 or new.casefold() in taken:
            # Never truncate through a token. A unique, restorable name instead.
            if reserved is None:
                reserved = ids.load_reserved()
            new = f"SHEET_{ids.new_id(lambda c: c in reserved)}"
        taken.discard(old.casefold())
        taken.add(new.casefold())
        renames[old] = new
    if renames:
        log.info("xlsx sheet titles anonymized: %d", len(renames))
    return renames


def _apply_sheet_renames(wb, renames: dict[str, str]) -> None:
    # Two phases so A->B, B->A style swaps can't collide mid-way.
    temp = {}
    for i, old in enumerate(renames):
        wb[old].title = f"__docanon_tmp_{i}"
        temp[f"__docanon_tmp_{i}"] = renames[old]
    for t, final in temp.items():
        wb[t].title = final


def _rewrite_sheet_refs(xml: str, renames: dict[str, str]) -> str:
    """Point formula / defined-name / pivot references at the renamed sheets."""
    def text_safe(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    for old, new in renames.items():
        # Inside a formula reference an apostrophe is written twice: 'O''Brien'!A1
        quoted_new = text_safe("'" + new.replace("'", "''") + "'!")
        oq = old.replace("'", "''")
        for o in {oq, text_safe(oq), xml_escape(oq)}:
            xml = xml.replace(f"'{o}'!", quoted_new)
            xml = xml.replace(f"&apos;{o}&apos;!", quoted_new)
        if re.fullmatch(r"[^\W\d][\w.]*", old):          # unquoted refs: Smith!A1, José!A1
            # Keep the reference's own style: unquoted stays unquoted when the
            # new name allows it (tokens do), so a restore is text-exact.
            bare_new = text_safe(new + "!") if re.fullmatch(r"[^\W\d][\w.]*", new) else quoted_new
            xml = re.sub(rf"(?<![\w.']){re.escape(old)}!", lambda _m: bare_new, xml)
        xml = xml.replace(f'sheet="{xml_escape(old)}"', f'sheet="{xml_escape(new)}"')
    return xml


# In worksheet XML, <v> holds numbers and shared-string indexes. Real cell
# text was already handled by openpyxl and the shared-strings pass, so the
# raw pass must never rewrite a <v>.
def _protect_cell_values(name: str) -> bool:
    return _XLSX_SHEET_PARTS.fullmatch(name) is not None


_NUMERIC_NODE_RE = re.compile(r"[\s\d.,:+\-]*")
_ATTR_RE = re.compile(r"""([\w:.-]+)\s*=\s*("[^"]*"|'[^']*')""")
_ELEMENT_RE = re.compile(r"</?\s*([\w:.-]+)")
# Attribute values that are coordinates, indexes, or identifiers. PII never
# legitimately lives there, and rewriting them breaks the file.
_STRUCTURAL_ATTRS = {"r", "ref", "sqref", "spans", "s", "t", "type", "count", "uniquecount"}


def _local(name: str) -> str:
    return name.split(":")[-1].lower()


def _markup_span_ok(text: str, lt: int, s0: int, e0: int) -> bool:
    """May the raw pass rewrite text[s0:e0], which sits inside the tag at lt?

    Only inside a non-structural attribute value (author, title, Target,
    name...), and only for values of 4+ characters that are not short
    numbers. Never element/attribute names, coordinates (r="B7"), ids, or
    Word revision ids (rsid).
    """
    matched = text[s0:e0]
    if len(matched) < 4 or is_short_number(matched):
        return False
    gt = text.find(">", e0)
    tag = text[lt:(gt + 1 if gt != -1 else len(text))]
    el = _ELEMENT_RE.match(tag)
    if el and _local(el.group(1)).startswith("rsid"):
        return False
    rs, re0 = s0 - lt, e0 - lt
    for m in _ATTR_RE.finditer(tag):
        if m.start(2) + 1 <= rs and re0 <= m.end(2) - 1:
            attr = _local(m.group(1))
            return not (attr in _STRUCTURAL_ATTRS or attr.startswith("rsid") or attr.endswith("id"))
    return False


def _apply_raw(text: str, replacer: Replacer, protect_v: bool = False) -> str:
    """Raw-XML substitution that never damages markup.

    By the time this runs, real document text has been handled by the
    format-aware passes. What is left is hidden layers (metadata, formula
    literals, alt text) - and markup. Rules:
      - inside a tag: only non-structural attribute values (_markup_span_ok)
      - xlsx <v> cell values: never (numbers and shared-string indexes)
      - formulas (<f>, <definedName>): only inside "string literals" - never
        a reference (B7, A94) or operand
      - number-only nodes (word counts, offsets): never a short number
    Letters and long numbers (phones, SSNs) are replaced everywhere else.
    """
    spans = replacer.find(text)
    if not spans:
        return text
    keep = []
    for s0, e0, r in spans:
        lt, gt = text.rfind("<", 0, s0), text.rfind(">", 0, s0)
        if lt > gt:                                   # inside a tag
            if not _markup_span_ok(text, lt, s0, e0):
                continue
        else:
            a = gt + 1
            open_lt = text.rfind("<", 0, a)
            el = _ELEMENT_RE.match(text, open_lt) if open_lt != -1 else None
            element = _local(el.group(1)) if el else ""
            if protect_v and element == "v":
                continue
            if element in ("f", "definedname"):
                # Formula syntax (cells and defined names). PII can only live
                # in a "string literal"; outside quotes everything is a
                # reference or operator (B7, A94, 2).
                before = text[a:s0].replace("&quot;", '"')
                in_string = before.count('"') % 2 == 1
                if not in_string:
                    if not getattr(replacer, "restoring", False):
                        continue
                    # Restoring a token that named a sheet: inside 'quoted'!
                    # references an apostrophe must be doubled (O''Brien).
                    r = r.replace("&apos;", "&apos;&apos;").replace("'", "''")
            if is_short_number(text[s0:e0]):
                b = text.find("<", e0)
                node = text[a:(b if b != -1 else len(text))]
                if element == "v" or _NUMERIC_NODE_RE.fullmatch(node):
                    continue
        keep.append((s0, e0, r))
    return splice(text, keep)


# ---------------------------------------------------------------------------
# Deep scrub primitives
# ---------------------------------------------------------------------------

def _deep_scrub_zip(
    src: Path,
    dst: Path,
    replacer: Replacer,
    per_part_handlers: dict | None = None,
    structured_replace=None,
    raw_protect=None,
) -> list[str]:
    """Walk every part in `src`, run handlers and the raw map substitution, repack to `dst`.

    `structured_replace(name, xml, map) -> xml` runs first on text parts - it
    is the run-aware pass that catches strings split across XML runs.
    `raw_protect(name) -> bool` marks parts whose <v> cell values are off-limits.
    Returns a list of layer labels for logging.
    """
    per_part_handlers = per_part_handlers or {}
    layers_touched: set[str] = set()
    # Raw XML variant: matches XML-escaped originals ("Smith &amp; Sons") and
    # writes XML-safe replacements.
    raw_replacer = replacer.for_raw_xml()

    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            handler = _find_handler(info.filename, per_part_handlers)

            # Drop sentinel - skip writing this part entirely.
            if handler is _drop_part:
                layers_touched.add(f"drop:{info.filename}")
                continue

            # 1. Structured (run-aware) replacement, then raw map substitution
            #    for any text / XML part. Apply before the metadata-zero
            #    handlers so any placeholder that lands in author/company
            #    fields is then cleared by the explicit handler.
            if _is_text_part(info.filename):
                try:
                    text = data.decode("utf-8")
                    new_text = _decode_char_refs(text)
                    if structured_replace is not None:
                        new_text = structured_replace(info.filename, new_text, replacer)
                        if new_text != text:
                            layers_touched.add("run_aware_substitution")
                    protect_v = bool(raw_protect and raw_protect(info.filename))
                    substituted = _apply_raw(new_text, raw_replacer, protect_v)
                    if substituted != new_text:
                        layers_touched.add("raw_xml_substitution")
                    if substituted != text:
                        data = substituted.encode("utf-8")
                except UnicodeDecodeError:
                    pass

            # 2. Targeted handler - tracked-change strip, metadata zero, etc.
            if handler is not None:
                try:
                    new_data, label = handler(info.filename, data)
                    if new_data is _DROP:
                        layers_touched.add(f"drop:{info.filename}")
                        continue
                    data = new_data
                    if label:
                        layers_touched.add(label)
                except Exception as exc:  # never let a malformed part take down the run
                    log.warning("handler failed on %s: %s", info.filename, exc)

            zout.writestr(info, data)

    return sorted(layers_touched)


_DROP = object()  # sentinel for handlers that signal "drop this part entirely"


def _find_handler(name: str, table: dict):
    for key, handler in table.items():
        if isinstance(key, str) and "*" in key:
            # cheap glob: replace * with non-slash run
            pat = re.compile(re.escape(key).replace(r"\*", r"[^/]*"))
            if pat.fullmatch(name):
                return handler
        elif isinstance(key, str):
            if key == name:
                return handler
        else:  # compiled pattern
            if key.fullmatch(name) or key.match(name):
                return handler
    return None


def _drop_part(name: str, data: bytes) -> tuple:  # pragma: no cover - sentinel target
    return _DROP, f"drop:{name}"


def _empty_comments_xml(name: str, data: bytes) -> tuple[bytes, str]:
    empty = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'
    return empty, "comments_cleared"


def _empty_comments_ex_xml(name: str, data: bytes) -> tuple[bytes, str]:
    # Emptied rather than dropped: dropping leaves stale [Content_Types].xml
    # and .rels references that can trigger Office repair prompts (M4).
    empty = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"/>'
    return empty, "comments_extended_cleared"


def _empty_xlsx_comments_xml(name: str, data: bytes) -> tuple[bytes, str]:
    empty = b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><authors/><commentList/></comments>'
    return empty, "xlsx_comments_cleared"


def _strip_revision_marks(name: str, data: bytes) -> tuple[bytes, str]:
    """Remove <w:ins>/<w:del> elements while keeping the final text.

    Strategy: drop opening/closing tags so the inner runs survive in place
    (this is a conservative XML-aware regex). Tracked-change deletions are
    fully removed (their contents shouldn't survive in the final document).
    """
    text = data.decode("utf-8", errors="ignore")

    # Remove <w:del>...</w:del> entirely (deleted content)
    text = re.sub(r"<w:del\b[^>]*>.*?</w:del>", "", text, flags=re.DOTALL)
    # Unwrap <w:ins>...</w:ins> -> keep inner content
    text = re.sub(r"<w:ins\b[^>]*>", "", text)
    text = re.sub(r"</w:ins>", "", text)
    # Strip alt text descriptions that may carry names ("Photo of John")
    text = re.sub(r' descr="[^"]*"', '', text)
    # Remove comment anchors - their comments part is emptied, so dangling
    # references would otherwise risk repair prompts (M4).
    text = re.sub(r"<w:commentRangeStart[^>]*/>", "", text)
    text = re.sub(r"<w:commentRangeEnd[^>]*/>", "", text)
    text = re.sub(r"<w:commentReference[^>]*/>", "", text)

    return text.encode("utf-8"), "tracked_changes_stripped"


def _zero_core_authors(name: str, data: bytes) -> tuple[bytes, str]:
    text = data.decode("utf-8", errors="ignore")
    # Empty the inner text. Preserve the open tag (with any xmlns/etc attributes)
    # so the document remains a valid OOXML core-properties file.
    text = re.sub(
        r"(<dc:creator(?:\s[^>]*)?>)[^<]*(</dc:creator>)",
        r"\1\2", text,
    )
    text = re.sub(
        r"(<cp:lastModifiedBy(?:\s[^>]*)?>)[^<]*(</cp:lastModifiedBy>)",
        r"\1\2", text,
    )
    return text.encode("utf-8"), "metadata_zeroed"


def _zero_app_company(name: str, data: bytes) -> tuple[bytes, str]:
    text = data.decode("utf-8", errors="ignore")
    text = re.sub(r"(<Company(?:\s[^>]*)?>)[^<]*(</Company>)", r"\1\2", text)
    text = re.sub(r"(<Manager(?:\s[^>]*)?>)[^<]*(</Manager>)", r"\1\2", text)
    return text.encode("utf-8"), "app_metadata_zeroed"


_TEXT_PART_SUFFIXES = (".xml", ".rels", ".txt", ".vml")


def _is_text_part(name: str) -> bool:
    n = name.lower()
    return any(n.endswith(s) for s in _TEXT_PART_SUFFIXES)
