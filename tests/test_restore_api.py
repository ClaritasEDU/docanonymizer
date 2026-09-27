"""Restore routes + the whole spreadsheet -> AI -> restore flow through the API."""

from __future__ import annotations

import io
import json
import re
import time
from unittest.mock import patch

JANE = "3A4F9C2B1D0E"


def _client():
    from app.server import create_app
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def _save_key(name, rmap, **extra):
    from app.config import KEYS_DIR
    payload = {"session_id": "abcd1234", "original_filename": "donors.xlsx",
               "replacement_map": rmap, **extra}
    (KEYS_DIR / name).write_text(json.dumps(payload))
    return name


def test_keys_listing_reports_format_and_counts():
    _save_key("new_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"}, id_format="hex12")
    _save_key("old_11111111.key.json", {"Jane Smith": "[PERSON_3A4F]", "Bob": "[PERSON_11C2]"})
    keys = {k["name"]: k for k in _client().get("/api/keys").get_json()["keys"]}
    assert keys["new_abcd1234.key.json"]["format"] == "hex12"
    assert keys["new_abcd1234.key.json"]["ids"] == 1
    assert keys["old_11111111.key.json"]["format"] == "legacy"
    assert keys["old_11111111.key.json"]["ids"] == 2
    assert keys["old_11111111.key.json"]["original_filename"] == "donors.xlsx"


def test_keys_listing_hides_the_ledger():
    from app import ids
    ids.LEDGER_FILE.write_text("ABCDEF123456\n")
    assert _client().get("/api/keys").get_json()["keys"] == []


def test_restore_text_with_saved_key():
    name = _save_key("d_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    r = _client().post("/api/unanonymize/text", json={
        "text": f"Top donor: person_{JANE.lower()} and [PERSON_ABCDEF123456]",
        "keys": [name],
    })
    assert r.status_code == 200
    body = r.get_json()
    assert body["text"] == "Top donor: Jane Smith and [PERSON_ABCDEF123456]"
    assert body["report"]["restored"] == 1
    assert body["report"]["unresolved"] == ["[PERSON_ABCDEF123456]"]


def test_restore_text_writes_nothing_to_disk():
    from app.config import OUTPUT_DIR, UPLOADS_DIR
    name = _save_key("d_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    _client().post("/api/unanonymize/text", json={"text": f"[PERSON_{JANE}]", "keys": [name]})
    assert list(OUTPUT_DIR.iterdir()) == [] and list(UPLOADS_DIR.iterdir()) == []


def test_restore_text_input_errors():
    c = _client()
    name = _save_key("d_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    assert c.post("/api/unanonymize/text", json={"text": "  ", "keys": [name]}).status_code == 400
    assert c.post("/api/unanonymize/text", json={"text": "x", "keys": []}).status_code == 400
    r = c.post("/api/unanonymize/text", json={"text": "x", "keys": ["nope.key.json"]})
    assert r.status_code == 400 and "not found" in r.get_json()["error"]
    for evil in ("../d_abcd1234.key.json", "/etc/passwd", "d_abcd1234.key.json/..", "x.json",
                 "..\\d_abcd1234.key.json", ".key.json", 42, None):
        assert c.post("/api/unanonymize/text", json={"text": "x", "keys": [evil]}).status_code == 400


def test_hand_copied_key_with_spaces_is_usable():
    """A key dragged into keys/ in Finder keeps names like "Donor List (1)"."""
    name = _save_key("Donor List (1).key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    c = _client()
    assert name in [k["name"] for k in c.get("/api/keys").get_json()["keys"]]
    r = c.post("/api/unanonymize/text", json={"text": f"[PERSON_{JANE}]", "keys": [name]})
    assert r.status_code == 200 and r.get_json()["text"] == "Jane Smith"


def test_restore_text_conflicting_keys_409():
    a = _save_key("a_11111111.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    b = _save_key("b_22222222.key.json", {"Bob Torres": f"[PERSON_{JANE}]"})
    r = _client().post("/api/unanonymize/text", json={"text": f"[PERSON_{JANE}]", "keys": [a, b]})
    assert r.status_code == 409
    assert "a_11111111.key.json" in r.get_json()["error"]


def test_restore_file_with_saved_keys_list():
    from app.config import UPLOADS_DIR
    name = _save_key("d_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    c = _client()
    r = c.post("/api/unanonymize", data={
        "file": (io.BytesIO(f"# Answer\n**PERSON_{JANE}** gave most".encode()), "answer.md"),
        "keys": json.dumps([name]),
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    body = r.get_json()
    assert body["output_filename"] == "answer_restored.md"
    assert body["text"] == "# Answer\n**Jane Smith** gave most"
    assert body["report"]["restored"] == 1 and body["report"]["residual"] == 0
    assert c.get(body["download_url"]).data == b"# Answer\n**Jane Smith** gave most"
    assert list(UPLOADS_DIR.iterdir()) == []   # temp upload cleaned


def test_restore_file_needs_a_key_and_a_supported_type():
    c = _client()
    r = c.post("/api/unanonymize", data={"file": (io.BytesIO(b"x"), "a.txt")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    name = _save_key("d_abcd1234.key.json", {"Jane Smith": f"[PERSON_{JANE}]"})
    r = c.post("/api/unanonymize", data={"file": (io.BytesIO(b"x"), "a.exe"),
                                         "keys": json.dumps([name])},
               content_type="multipart/form-data")
    assert r.status_code == 400


def test_import_key():
    from app.config import KEYS_DIR, UPLOADS_DIR
    c = _client()
    payload = {"session_id": "x", "original_filename": "a.xlsx",
               "replacement_map": {"Jane Smith": f"[PERSON_{JANE}]"}}
    r = c.post("/api/keys/import", data={
        "key": (io.BytesIO(json.dumps(payload).encode()), "from laptop.key.json"),
    }, content_type="multipart/form-data")
    assert r.status_code == 200
    name = r.get_json()["imported"]
    assert name == "from_laptop.key.json" and (KEYS_DIR / name).exists()
    # Importing again never overwrites.
    r = c.post("/api/keys/import", data={
        "key": (io.BytesIO(json.dumps(payload).encode()), "from laptop.key.json"),
    }, content_type="multipart/form-data")
    assert r.get_json()["imported"] == "from_laptop_1.key.json"
    # Junk is refused and not left behind.
    r = c.post("/api/keys/import", data={"key": (io.BytesIO(b"{nope"), "bad.json")},
               content_type="multipart/form-data")
    assert r.status_code == 400
    assert list(UPLOADS_DIR.iterdir()) == []


def test_imported_key_ids_are_reserved_for_new_runs():
    from app import ids
    c = _client()
    payload = {"session_id": "x", "original_filename": "a.xlsx",
               "replacement_map": {"Jane Smith": "[PERSON_ABCDEF123456]"}}
    c.post("/api/keys/import", data={
        "key": (io.BytesIO(json.dumps(payload).encode()), "a.key.json"),
    }, content_type="multipart/form-data")
    assert "ABCDEF123456" in ids.load_reserved()


def _run_anonymize(c, data, filename, fake_items):
    from app import detector
    with patch.object(detector.llm, "llm_call", return_value=json.dumps(fake_items)):
        r = c.post("/api/anonymize/upload", data={"file": (data, filename), "tags": "ALL"},
                   content_type="multipart/form-data")
        sid = r.get_json()["session_id"]
        for _ in range(200):
            s = c.get(f"/api/anonymize/{sid}/status").get_json()
            if s.get("detection_complete") or s.get("error"):
                break
            time.sleep(0.05)
        assert s.get("detection_complete"), s
        c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
        for _ in range(200):
            res = c.get(f"/api/anonymize/{sid}/results").get_json()
            if res.get("verify_result") or res.get("error"):
                break
            time.sleep(0.05)
    assert res["verify_result"]["passed"], res
    return sid, s["preview"], res


def test_spreadsheet_to_ai_and_back_through_the_api():
    """The operator's flow, end to end:
    add a spreadsheet -> names/emails/phones/addresses become unique 12-char
    IDs -> view the output -> an AI answers with mangled IDs -> restore."""
    from openpyxl import Workbook

    people = [
        ("Jane Smith", "jane@parish.org", "(512) 555-0101", "12 Oak St, Austin TX 78701"),
        ("Bob Torres", "bob@parish.org", "(512) 555-0102", "9 Elm Ave, Austin TX 78702"),
        ("Ann Lee", "ann@parish.org", "(512) 555-0103", "4 Pine Rd, Austin TX 78703"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Email", "Phone", "Address", "Gift"])
    for i, p in enumerate(people):
        ws.append(list(p) + [100 * (i + 1)])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fake = [{"text": v, "type": t} for p in people
            for v, t in zip(p, ("PERSON", "EMAIL", "PHONE", "ADDRESS"))]

    c = _client()
    sid, preview, res = _run_anonymize(c, buf, "donors.xlsx", fake)
    assert preview["entities"] == 12

    # "Show me the output": every value replaced, every ID unique.
    text = c.get(f"/api/anonymize/{sid}/text").get_json()["text"]
    placeholders = re.findall(r"\[([A-Z]+)_([0-9A-F]{12})\]", text)
    assert len(placeholders) == 12
    assert len({h for _, h in placeholders}) == 12
    for p in people:
        for v in p:
            assert v not in text

    # The AI's answer: ranks donors using mangled IDs.
    rows = [line.split("\t") for line in text.strip().splitlines()[1:]]
    answer = "\n".join(
        f"{rank}. {name.lower()} - email {email.strip('[]')}, call {phone[1:-1]} ({gift})"
        for rank, (name, email, phone, _addr, gift) in enumerate(reversed(rows), start=1)
    )

    key_name = res["key_filename"]
    listed = [k["name"] for k in c.get("/api/keys").get_json()["keys"]]
    assert key_name in listed
    r = c.post("/api/unanonymize/text", json={"text": answer, "keys": listed})
    assert r.status_code == 200
    body = r.get_json()
    expected = "\n".join(
        f"{rank}. {n} - email {e}, call {ph} ({100 * (3 - rank + 1)})"
        for rank, (n, e, ph, _a) in enumerate(reversed(people), start=1)
    )
    assert body["text"] == expected
    assert body["report"]["restored"] == 9
    assert body["report"]["unresolved_count"] == 0


def test_two_spreadsheets_never_share_an_id():
    """Anonymize two files one after another - no ID appears in both keys,
    so pasting both into one AI chat can't cause a mix-up."""
    from app.config import KEYS_DIR
    c = _client()
    fake_a = [{"text": f"Person A{i}", "type": "PERSON"} for i in range(40)]
    fake_b = [{"text": f"Person B{i}", "type": "PERSON"} for i in range(40)]
    body_a = "\n".join(i["text"] for i in fake_a).encode()
    body_b = "\n".join(i["text"] for i in fake_b).encode()
    _, _, ra = _run_anonymize(c, io.BytesIO(body_a), "a.txt", fake_a)
    _, _, rb = _run_anonymize(c, io.BytesIO(body_b), "b.txt", fake_b)
    ka = json.loads((KEYS_DIR / ra["key_filename"]).read_text())
    kb = json.loads((KEYS_DIR / rb["key_filename"]).read_text())
    assert ka["id_format"] == kb["id_format"] == "hex12"
    assert not set(ka["entity_registry"]) & set(kb["entity_registry"])
    assert len(ka["entity_registry"]) == len(kb["entity_registry"]) == 40



def test_restore_text_rejects_non_object_json():
    c = _client()
    for body in ([1], "text", 5):
        assert c.post("/api/unanonymize/text", json=body).status_code == 400


def test_preview_type_toggle_keeps_whole_type_as_text():
    """[x] ORG toggled off in the preview keeps every ORG value as original
    text across the whole document - including past the preview window."""
    from app import detector
    c = _client()
    body = ("Jane Smith of Acme Corp. " + "x" * 11000 + " Acme Corp again, and Jane Smith.").encode()
    fake = json.dumps([{"text": "Jane Smith", "type": "PERSON"}, {"text": "Acme Corp", "type": "ORG"},
                       {"text": "250", "type": "FINANCIAL"}])
    with patch.object(detector.llm, "llm_call", return_value=fake):
        sid = c.post("/api/anonymize/upload", data={"file": (io.BytesIO(body), "memo.txt"), "tags": "ALL"},
                     content_type="multipart/form-data").get_json()["session_id"]
        for _ in range(100):
            s = c.get(f"/api/anonymize/{sid}/status").get_json()
            if s.get("detection_complete"):
                break
            time.sleep(0.05)
        assert s["preview"]["kept_amounts"] == 1
        c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": [], "deselected_types": ["ORG", "BOGUS"]})
        for _ in range(100):
            res = c.get(f"/api/anonymize/{sid}/results").get_json()
            if res.get("verify_result"):
                break
            time.sleep(0.05)
    assert res["verify_result"]["passed"]
    text = c.get(f"/api/anonymize/{sid}/text").get_json()["text"]
    assert text.count("Acme Corp") == 2 and "Jane Smith" not in text
