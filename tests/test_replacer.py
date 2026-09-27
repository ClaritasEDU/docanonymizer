"""Single-pass replacement engine (replacer.py) and the verifier's use of it."""

from __future__ import annotations

import random
import time


def _brute(text, mapping):
    """Reference scanner: letter-bearing originals leftmost-longest, then
    short numbers in the remaining gaps, whole numbers only."""
    from app.replacer import _number_boundary_ok, is_short_number
    texty = sorted((w for w in mapping if not is_short_number(w)), key=len, reverse=True)
    numbers = sorted((w for w in mapping if is_short_number(w)), key=len, reverse=True)
    out, used, i = [], [False] * len(text), 0
    while i < len(text):
        for w in texty:
            if text.startswith(w, i):
                out.append((i, i + len(w), mapping[w]))
                for k in range(i, i + len(w)):
                    used[k] = True
                i += len(w)
                break
        else:
            i += 1
    for i in range(len(text)):
        if used[i]:
            continue
        for w in numbers:
            e = i + len(w)
            if (text.startswith(w, i) and not any(used[i:e])
                    and _number_boundary_ok(text, i, e)):
                out.append((i, e, mapping[w]))
                for k in range(i, e):
                    used[k] = True
                break
    return sorted(out)


def test_matches_reference_on_random_inputs():
    from app.replacer import LiteralReplacer
    rng = random.Random(11)
    alpha = "abAB 1."
    for _ in range(1500):
        words = {"".join(rng.choice(alpha) for _ in range(rng.randint(1, 5)))
                 for _ in range(rng.randint(1, 8))}
        mapping = {w: f"<{i}>" for i, w in enumerate(words)}
        text = "".join(rng.choice(alpha) for _ in range(rng.randint(0, 40)))
        assert LiteralReplacer(mapping, protect_placeholders=False).find(text) == _brute(text, mapping)


def test_fallback_scan_matches_trie():
    from app.replacer import LiteralReplacer
    rng = random.Random(5)
    alpha = "xyXY 2"
    for _ in range(500):
        words = {"".join(rng.choice(alpha) for _ in range(rng.randint(1, 4)))
                 for _ in range(rng.randint(1, 6))}
        mapping = {w: f"<{i}>" for i, w in enumerate(words)}
        text = "".join(rng.choice(alpha) for _ in range(rng.randint(0, 30)))
        fast = LiteralReplacer(mapping)
        slow = LiteralReplacer(mapping)
        slow._fallback, slow._words = True, sorted(mapping, key=len, reverse=True)
        assert fast.find(text) == slow.find(text)


def test_longest_wins_at_a_position():
    from app.replacer import LiteralReplacer, apply
    r = LiteralReplacer({"Jane": "[P1]", "Jane Smith": "[P2]"})
    assert apply("Jane Smith and Jane", r) == "[P2] and [P1]"


def test_short_original_never_rewrites_inside_a_placeholder():
    """The old sequential loop turned [PERSON_3A4F9C2B1D0E] into
    [PERSON_3A4F9C[SID_...]0E] when "2B1D" was also a value."""
    from app.replacer import LiteralReplacer, apply
    r = LiteralReplacer({
        "Jane Smith": "[PERSON_3A4F9C2B1D0E]",
        "2B1D": "[SID_AAAAAAAAAAA1]",
        "SON": "[ORG_111111111ABC]",
    })
    once = apply("Jane Smith 2B1D", r)
    assert once == "[PERSON_3A4F9C2B1D0E] [SID_AAAAAAAAAAA1]"
    assert apply(once, r) == once   # idempotent


def test_scrub_text_does_not_corrupt_placeholders(tmp_path):
    from app.extractors import extract
    from app.scrubber import scrub_text
    src = tmp_path / "m.txt"
    src.write_text("Jane Smith scored 94 on test 2B1D")
    rmap = {"Jane Smith": "[PERSON_3A4F9C2B1D0E]", "94": "[GRADE_5C94AB01DE2F]", "2B1D": "[SID_A2B1D0000001]"}
    out = tmp_path / "o.txt"
    scrub_text(extract(src), rmap, out)
    assert out.read_text() == "[PERSON_3A4F9C2B1D0E] scored [GRADE_5C94AB01DE2F] on test [SID_A2B1D0000001]"


def test_verifier_ignores_originals_inside_placeholders(tmp_path):
    """A 2-digit grade occurs inside ~4% of random 12-char IDs. That must not
    fail verification of a clean file."""
    from app.verifier import verify_output
    out = tmp_path / "o.txt"
    out.write_text("[PERSON_3A94F2B1C0DE] scored [GRADE_5C0000AB1DE2]")
    result = verify_output(out, {"Jane Smith": "[PERSON_3A94F2B1C0DE]", "94": "[GRADE_5C0000AB1DE2]"})
    assert result.passed


def test_verifier_still_catches_real_residue(tmp_path):
    from app.verifier import verify_output
    out = tmp_path / "o.txt"
    out.write_text("[PERSON_3A94F2B1C0DE] scored 94")
    result = verify_output(out, {"Jane Smith": "[PERSON_3A94F2B1C0DE]", "94": "[GRADE_5C0000AB1DE2]"})
    assert not result.passed
    assert result.map_match_types == ["GRADE"]


def test_xlsx_scrub_scales_linearly(tmp_path):
    """3,000 rows x 4 PII columns. The old loop was quadratic (~30s here)."""
    from openpyxl import Workbook, load_workbook
    from app.mapper import EntityRegistry
    from app.scrubber import scrub_xlsx
    wb = Workbook()
    ws = wb.active
    reg = EntityRegistry()
    for i in range(3000):
        row = [f"Person{i} Last{i}", f"p{i}@x.org", f"(512) 555-{i:04d}", f"{i} Main St"]
        ws.append(row + [f"Called Person{i} Last{i}"])
        for v, tag in zip(row, ("PERSON", "EMAIL", "PHONE", "ADDRESS")):
            reg.add(v, tag)
    src = tmp_path / "big.xlsx"
    wb.save(src)
    t = time.monotonic()
    scrub_xlsx(src, reg.as_replacement_map(), tmp_path / "out.xlsx")
    assert time.monotonic() - t < 15
    cells = [c.value for row in load_workbook(tmp_path / "out.xlsx").active.iter_rows() for c in row]
    assert cells[0] == reg.replacements["Person0 Last0"]
    assert cells[4] == f"Called {reg.replacements['Person0 Last0']}"


def test_short_numbers_match_whole_numbers_only():
    """A grade of 94 is not the 94 inside 1945, 94.5, 1,945, or cell A94's row."""
    from app.replacer import LiteralReplacer, apply
    r = LiteralReplacer({"94": "[G]", "3.87": "[P]"})
    assert apply("scored 94; gave 1945; 94.5 avg; 1,945 total; 194; 94", r) == \
        "scored [G]; gave 1945; 94.5 avg; 1,945 total; 194; [G]"
    assert apply("GPA 3.87 vs 13.87 vs 3.875", r) == "GPA [P] vs 13.87 vs 3.875"
    assert apply("Grade94 (94) 94th A94 $B$94 94%", r) == "Grade94 ([G]) 94th A94 $B$94 [G]%"


def test_comma_lists_are_lists_not_thousands():
    """204518,78704 is two values; 1,945 is one number (review finding 2)."""
    from app.replacer import LiteralReplacer, apply
    r = LiteralReplacer({"204518": "[S1]", "204519": "[S2]", "78704": "[Z]", "94": "[G]",
                         "87": "[H]", "945": "[X]"})
    assert apply("Jane,204518,78704\nSibling IDs: 204518,204519", r) == \
        "Jane,[S1],[Z]\nSibling IDs: [S1],[S2]"
    assert apply("scores 94,87,100.", r) == "scores [G],[H],100."
    assert apply("gift $1,945 and 1,945.50 and ip 10.0.94.1", r) == \
        "gift $1,945 and 1,945.50 and ip 10.0.94.1"


def test_pii_shaped_like_a_token_is_still_replaced_and_verified(tmp_path):
    """Protection covers only this map's own placeholders (review finding 1)."""
    from app.replacer import LiteralReplacer, apply
    from app.verifier import verify_output
    rmap = {"MRN_000123456789": "[ID_A1B2C3D4E5F6]", "Jane Smith": "[PERSON_3A4F9C2B1D0E]"}
    assert apply("Patient Jane Smith, MRN_000123456789", LiteralReplacer(rmap)) == \
        "Patient [PERSON_3A4F9C2B1D0E], [ID_A1B2C3D4E5F6]"
    leaky = tmp_path / "o.txt"
    leaky.write_text("Patient [PERSON_3A4F9C2B1D0E], MRN_000123456789")
    assert not verify_output(leaky, rmap).passed


def test_long_numbers_still_match_inside_longer_forms():
    """A phone is caught even inside +1 / extension forms (substring rule)."""
    from app.replacer import LiteralReplacer, apply
    r = LiteralReplacer({"5125550101": "[PHONE_A]"})
    assert apply("+15125550101 or 5125550101x2", r) == "+1[PHONE_A] or [PHONE_A]x2"


def test_is_short_number():
    from app.replacer import is_short_number
    assert is_short_number("94") and is_short_number("3.87") and is_short_number("78701")
    assert not is_short_number("(512)")            # self-delimited by the parens
    assert not is_short_number("5125550101")       # 10 digits
    assert not is_short_number("A+") and not is_short_number("Apt 4")
    assert not is_short_number("")


def test_raw_xml_pass_leaves_coordinates_ids_and_formulas_alone():
    """Review finding 6 + web agent finding: B7 must not hit <c r="B7">, A1 must
    not hit an rsid, and a grade 94 must not rewrite the formula A94*2."""
    from app.replacer import LiteralReplacer
    from app.scrubber import _apply_raw
    r = LiteralReplacer({"B7": "[ID_B7B7B7B7B7B7]", "A1D8": "[ID_A1D8A1D8A1D8]",
                         "94": "[GRADE_5C94AB01DE2F]", "Jane Smith": "[PERSON_3A4F9C2B1D0E]"})
    xml = ('<c r="B7" s="3"><f>A94*2</f></c><w:rsid w:val="00A1D8D"/>'
           '<w:p w:rsidR="00A1D8D2"/><Relationship Id="A1D8" Target="x"/>'
           '<w:comment w:author="Jane Smith"/><f>"Jane Smith"&amp;B7</f>')
    assert _apply_raw(xml, r) == (
        '<c r="B7" s="3"><f>A94*2</f></c><w:rsid w:val="00A1D8D"/>'
        '<w:p w:rsidR="00A1D8D2"/><Relationship Id="A1D8" Target="x"/>'
        '<w:comment w:author="[PERSON_3A4F9C2B1D0E]"/><f>"[PERSON_3A4F9C2B1D0E]"&amp;B7</f>')


def test_raw_xml_pass_never_writes_short_numbers_into_markup():
    from app.replacer import LiteralReplacer
    from app.scrubber import _apply_raw
    r = LiteralReplacer({"94": "[GRADE_5C94AB01DE2F]", "5125550101": "[PHONE_7C1B0A94E2D3]",
                         "Jane": "[PERSON_3A4F9C2B1D0E]"})
    xml = ('<row r="94"><c r="A94" s="94"><v>94</v></c></row>'
           '<Words>94</Words><t>score 94</t><a href="tel:5125550101" title="Jane"/>')
    assert _apply_raw(xml, r, protect_v=True) == (
        '<row r="94"><c r="A94" s="94"><v>94</v></c></row>'
        '<Words>94</Words><t>score [GRADE_5C94AB01DE2F]</t>'
        '<a href="tel:[PHONE_7C1B0A94E2D3]" title="[PERSON_3A4F9C2B1D0E]"/>')
