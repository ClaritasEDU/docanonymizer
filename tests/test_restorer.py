"""Restoring AI output (restorer.py + unanonymize.restore_file).

The promise under test: whatever an AI tool does to the identifiers - drop
the brackets, change case, markdown-escape them, relabel them, keep only the
hex - every identifier comes back as exactly the value it stood for, and
nothing is guessed when an identifier can't be matched.
"""

from __future__ import annotations

import random
import re

import pytest

JANE, BOB, EMAIL, ADDR = "3A4F9C2B1D0E", "11C2D3E4F5A6", "7C1B0A94E2D3", "0D1E2F3A4B5C"

KEY = {
    "session_id": "abcd1234",
    "original_filename": "donors.xlsx",
    "id_format": "hex12",
    "replacement_map": {
        "Jane Smith": f"[PERSON_{JANE}]",
        "Bob Torres": f"[PERSON_{BOB}]",
        "jane@x.org": f"[EMAIL_{EMAIL}]",
        "12 Oak St, Austin TX": f"[ADDRESS_{ADDR}]",
    },
}


def _index(*payloads):
    from app.restorer import build_index
    return build_index([(f"k{i}.key.json", p) for i, p in enumerate(payloads)])


def _restore(text, *payloads):
    from app.restorer import restore_text
    return restore_text(text, _index(*(payloads or (KEY,))))


@pytest.mark.parametrize("token", [
    f"[PERSON_{JANE}]",                 # exact
    f"[person_{JANE.lower()}]",         # case changed
    f"\\[PERSON\\_{JANE}\\]",           # markdown-escaped
    f"\\[PERSON_{JANE}\\]",
    f"[PERSON\\_{JANE}]",
    f"PERSON_{JANE}",                   # brackets dropped
    f"PERSON\\_{JANE}",
    f"[ PERSON_{JANE} ]",               # padded
    f"[DONOR_{JANE}]",                  # relabeled by the AI
    f"Donor_Name_{JANE}",
    JANE,                               # bare ID
    JANE.lower(),
])
def test_every_mangled_form_restores(token):
    out, rep = _restore(f"Top donor: {token}.")
    assert out == "Top donor: Jane Smith."
    assert rep.restored == 1 and rep.unresolved_count == 0
    assert rep.by_tag == {"PERSON": 1}


def test_relabel_and_untagged_are_counted():
    _, rep = _restore(f"[DONOR_{JANE}] and {BOB} and [EMAIL_{EMAIL}]")
    assert rep.restored == 3
    assert rep.relabeled == 1
    assert rep.untagged == 1


def test_markdown_table_and_json_contexts():
    text = (f"| name | email |\n|---|---|\n| **[PERSON_{JANE}]** | [EMAIL_{EMAIL}] |\n"
            f'{{"donor": "[PERSON_{BOB}]", "home": "ADDRESS_{ADDR}"}}')
    out, rep = _restore(text)
    assert out == ("| name | email |\n|---|---|\n| **Jane Smith** | jane@x.org |\n"
                   '{"donor": "Bob Torres", "home": "12 Oak St, Austin TX"}')
    assert rep.restored == 4


def test_list_brackets_are_not_eaten():
    """Brackets are consumed only as a pair around ONE token."""
    out, _ = _restore(f"[PERSON_{JANE}, PERSON_{BOB}]")
    assert out == "[Jane Smith, Bob Torres]"


def test_possessive_and_punctuation():
    out, _ = _restore(f"PERSON_{JANE}'s gift; ({BOB})")
    assert out == "Jane Smith's gift; (Bob Torres)"


def test_unresolved_tokens_are_reported_and_left_alone():
    truncated = f"[PERSON_{JANE[:-1]}]"            # AI dropped a character
    unknown = "[PERSON_ABCDEF123456]"              # not in the selected keys
    unknown_bare = "EMAIL_ABCDEF654321"
    text = f"{truncated} / {unknown} / {unknown_bare} / [PERSON_{JANE}]"
    out, rep = _restore(text)
    assert out == f"{truncated} / {unknown} / {unknown_bare} / Jane Smith"
    assert rep.restored == 1
    assert rep.unresolved == [truncated, unknown, unknown_bare]
    assert rep.unresolved_count == 3


def test_ordinary_text_is_not_touched_or_reported():
    text = ("file_2024 ID_2024 version 3A4F sha DEADBEEF1234 uuid "
            "123e4567-e89b-12d3-a456-426614174000 phone 5125550198 [NOTE_ABCD]")
    out, rep = _restore(text)
    assert out == text
    assert rep.restored == 0 and rep.unresolved_count == 0


def test_no_mixups_under_random_mangling():
    """1,000 distinct values, each mentioned several times in random mangled
    forms, interleaved: every mention restores to exactly its own value."""
    from app import ids
    from app.restorer import build_index, restore_text
    rng = random.Random(3)
    rmap, truth = {}, []
    for i in range(1000):
        tag = rng.choice(["PERSON", "EMAIL", "PHONE", "ADDRESS"])
        hx = ids.new_id()
        original = f"{tag.lower()}-value-{i} (#{rng.randint(0, 99)})"
        rmap[original] = f"[{tag}_{hx}]"
        truth.append((original, tag, hx))
    forms = [
        lambda t, h: f"[{t}_{h}]", lambda t, h: f"[{t.lower()}_{h.lower()}]",
        lambda t, h: f"\\[{t}\\_{h}\\]", lambda t, h: f"{t}_{h}",
        lambda t, h: f"[RELABELED_{h}]", lambda t, h: h,
    ]
    anon_parts, expected_parts = [], []
    for _ in range(4000):
        original, tag, hx = rng.choice(truth)
        anon_parts.append(rng.choice(forms)(tag, hx))
        expected_parts.append(original)
    sep = " | "
    index = build_index([("k.key.json", {"replacement_map": rmap})])
    out, rep = restore_text(sep.join(anon_parts), index)
    assert out == sep.join(expected_parts)
    assert rep.restored == 4000 and rep.unresolved_count == 0


# ---------------------------------------------------------------------------
# Keys: multiple, conflicting, legacy
# ---------------------------------------------------------------------------

def test_two_keys_restore_together():
    other = {"replacement_map": {"Ann Lee": "[PERSON_ABC123DEF456]"}}
    out, rep = _restore(f"[PERSON_{JANE}] met [PERSON_ABC123DEF456]", KEY, other)
    assert out == "Jane Smith met Ann Lee"
    assert len(rep.keys_used) == 2


def test_same_value_different_ids_across_keys_is_fine():
    other = {"replacement_map": {"Jane Smith": "[PERSON_ABC123DEF456]"}}
    out, _ = _restore(f"[PERSON_{JANE}] = [PERSON_ABC123DEF456]", KEY, other)
    assert out == "Jane Smith = Jane Smith"


def test_conflicting_keys_refuse_to_restore():
    from app.restorer import KeyConflictError
    bad = {"replacement_map": {"Somebody Else": f"[PERSON_{JANE}]"}}
    with pytest.raises(KeyConflictError, match="k0.key.json and k1.key.json"):
        _index(KEY, bad)


def test_legacy_key_needs_the_tag():
    legacy = {"replacement_map": {"Jane Smith": "[PERSON_3A4F]", "jane@x.org": "[EMAIL_3A4F]"}}
    out, rep = _restore("[PERSON_3A4F] / [email_3a4f] / PERSON_3A4F / 3A4F / [ORG_3A4F]", legacy)
    assert out == "Jane Smith / jane@x.org / Jane Smith / 3A4F / [ORG_3A4F]"
    assert rep.restored == 3
    assert rep.unresolved == ["[ORG_3A4F]"]


def test_legacy_shared_placeholder_restores_longest_and_warns():
    """v1.3 could give "Jane Smith" and "Jane" the same placeholder - the
    exact mix-up the 12-char rule removes. Old keys still restore, visibly."""
    legacy = {"replacement_map": {"Jane": "[PERSON_3A4F]", "Jane Smith": "[PERSON_3A4F]"}}
    out, rep = _restore("[PERSON_3A4F]", legacy)
    assert out == "Jane Smith"
    assert rep.legacy_ambiguous == 1


def test_legacy_conflict_across_keys_refuses():
    from app.restorer import KeyConflictError
    a = {"replacement_map": {"Jane Smith": "[PERSON_3A4F]"}}
    b = {"replacement_map": {"Bob Torres": "[PERSON_3A4F]"}}
    with pytest.raises(KeyConflictError):
        _index(a, b)


def test_empty_index_restores_nothing():
    out, rep = _restore(f"[PERSON_{JANE}]", {"replacement_map": {}})
    assert out == f"[PERSON_{JANE}]" and rep.restored == 0


# ---------------------------------------------------------------------------
# Files
# ---------------------------------------------------------------------------

def test_restore_markdown_file(tmp_path):
    from app.unanonymize import restore_file
    src = tmp_path / "analysis.md"
    src.write_text(f"# Findings\n- **person_{JANE.lower()}** gave the most\n- Contact: {EMAIL}\n")
    res = restore_file(src, _index(KEY), display_name="analysis.md")
    assert res.output_path.name == "analysis_restored.md"
    assert res.output_path.read_text() == "# Findings\n- **Jane Smith** gave the most\n- Contact: jane@x.org\n"
    assert res.report.restored == 2 and res.report.residual == 0


def test_restore_csv_from_ai(tmp_path):
    import csv
    from app.unanonymize import restore_file
    src = tmp_path / "ai_summary.csv"
    src.write_text(f"donor,total\n[PERSON_{JANE}],500\nPERSON_{BOB},250\n")
    res = restore_file(src, _index(KEY))
    rows = list(csv.reader(res.output_path.open()))
    assert rows == [["donor", "total"], ["Jane Smith", "500"], ["Bob Torres", "250"]]


def test_restore_xlsx_the_ai_rebuilt(tmp_path):
    """AI returns its own workbook: new columns, lowercased tokens, a total row."""
    from openpyxl import Workbook, load_workbook
    from app.unanonymize import restore_file
    wb = Workbook()
    ws = wb.active
    ws.append(["Donor", "Email", "Segment", "Total"])
    ws.append([f"[person_{JANE.lower()}]", f"EMAIL_{EMAIL}", "major", 500])
    ws.append([f"[PERSON_{BOB}]", "", "lapsed", 20])
    ws.append(["TOTAL", "", "", "=SUM(D2:D3)"])
    src = tmp_path / "segments_anon_ab12cd34.xlsx"
    wb.save(src)
    res = restore_file(src, _index(KEY), display_name=src.name)
    assert res.output_path.name == "segments_restored.xlsx"
    ws2 = load_workbook(res.output_path).active
    assert [c.value for c in ws2[2]] == ["Jane Smith", "jane@x.org", "major", 500]
    assert ws2["A3"].value == "Bob Torres"
    assert ws2["D4"].value == "=SUM(D2:D3)"
    assert res.report.restored == 3 and res.report.residual == 0


def test_restore_docx_split_runs_and_xml_escaping(tmp_path):
    """Token split across runs + originals containing & and < must come back
    intact and leave a valid document. Comments/metadata are not stripped."""
    from docx import Document
    from app.unanonymize import restore_file
    key = {"replacement_map": {
        "Smith & Sons <Ltd>": "[ORG_A1B2C3D4E5F6]",
        "Jane Smith": f"[PERSON_{JANE}]",
    }}
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Vendor: [ORG_A1B2C3")
    p.add_run("D4E5F6], contact ").bold = True
    p.add_run(f"PERSON_{JANE}")
    doc.core_properties.author = "Analyst"
    src = tmp_path / "memo.docx"
    doc.save(src)
    res = restore_file(src, _index(key))
    doc2 = Document(res.output_path)
    assert doc2.paragraphs[0].text == "Vendor: Smith & Sons <Ltd>, contact Jane Smith"
    assert doc2.core_properties.author == "Analyst"
    assert res.report.residual == 0


def test_full_round_trip_through_a_mangling_ai(tmp_path):
    """Anonymize a spreadsheet, let an 'AI' rewrite it with mangled tokens,
    restore - every value lands back on its own row."""
    from openpyxl import Workbook
    from app.extractors import extract
    from app.mapper import EntityRegistry
    from app.restorer import build_index, restore_text
    from app.scrubber import scrub_xlsx

    people = [(f"Person {i} Doe", f"p{i}@parish.org", f"(512) 555-{i:04d}", f"{i} Elm St")
              for i in range(200)]
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "email", "phone", "address"])
    reg = EntityRegistry()
    for row in people:
        ws.append(list(row))
        for v, tag in zip(row, ("PERSON", "EMAIL", "PHONE", "ADDRESS")):
            reg.add(v, tag)
    src = tmp_path / "roster.xlsx"
    wb.save(src)
    out = tmp_path / "roster_anon.xlsx"
    scrub_xlsx(src, reg.as_replacement_map(), out)
    anon_lines = extract(out).text.strip().splitlines()[1:]
    assert len(anon_lines) == 200

    # The "AI": reorders rows, drops brackets on emails, lowercases phones,
    # relabels addresses, and writes prose.
    rng = random.Random(9)
    order = list(range(200))
    rng.shuffle(order)
    answer, expected = [], []
    for i in order:
        name, email, phone, addr = anon_lines[i].split("\t")
        answer.append(f"- {name} ({email.strip('[]')}, {phone.lower()}) lives at "
                      f"{addr.replace('ADDRESS', 'HOME')}")
        n, e, ph, a = people[i]
        expected.append(f"- {n} ({e}, {ph}) lives at {a}")
    index = build_index([("roster.key.json", {"replacement_map": reg.as_replacement_map()})])
    restored, rep = restore_text("\n".join(answer), index)
    assert restored == "\n".join(expected)
    assert rep.restored == 800 and rep.relabeled == 200 and rep.unresolved_count == 0


def test_unanonymize_file_single_key_still_works(tmp_path):
    """Backwards-compatible entry point used by older callers."""
    from app.unanonymize import unanonymize_file
    src = tmp_path / "x.txt"
    src.write_text(f"[PERSON_{JANE}]")
    assert unanonymize_file(src, KEY).read_text() == "Jane Smith"


def test_restore_logs_contain_no_values(tmp_path):
    from app.config import LOG_FILE
    from app.unanonymize import restore_file
    src = tmp_path / "a.txt"
    src.write_text(f"[PERSON_{JANE}] at [ADDRESS_{ADDR}]")
    restore_file(src, _index(KEY))
    _restore(f"[PERSON_{BOB}] [EMAIL_{EMAIL}]")
    log = LOG_FILE.read_text()
    for value in KEY["replacement_map"]:
        assert value not in log
    assert re.search(r"restored=\d+", log)


# ---------------------------------------------------------------------------
# Real-world spreadsheet shapes
# ---------------------------------------------------------------------------

def test_numeric_phone_zip_and_date_cells_are_scrubbed(tmp_path):
    """Phones typed without dashes and ZIPs are stored as numbers, birthdays
    as dates. They used to survive the cell pass and get a placeholder jammed
    into a numeric cell by the raw pass - a workbook Excel had to repair."""
    import datetime
    from openpyxl import Workbook, load_workbook
    from app.extractors import extract
    from app.mapper import EntityRegistry
    from app.scrubber import scrub_xlsx
    from app.verifier import verify_output
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Phone", "Zip", "Born", "Gift"])
    ws.append(["Jane Smith", 5125550101, 78701, datetime.datetime(2011, 3, 14), 250])
    src = tmp_path / "n.xlsx"
    wb.save(src)
    text = extract(src).text                     # what the model sees
    assert "5125550101" in text and "2011-03-14 00:00:00" in text
    reg = EntityRegistry()
    for v, t in (("Jane Smith", "PERSON"), ("5125550101", "PHONE"),
                 ("78701", "ADDRESS"), ("2011-03-14", "DOB")):
        reg.add(v, t)
    out = tmp_path / "o.xlsx"
    scrub_xlsx(src, reg.as_replacement_map(), out)
    row = [c.value for c in load_workbook(out).active[2]]
    assert row[:3] == [reg.replacements["Jane Smith"], reg.replacements["5125550101"],
                       reg.replacements["78701"]]
    assert row[3] == f"{reg.replacements['2011-03-14']} 00:00:00"
    assert row[4] == 250                          # untouched number stays a number
    assert verify_output(out, reg.as_replacement_map()).passed


def test_short_value_never_corrupts_shared_string_indexes(tmp_path):
    """A grade "94" must not rewrite <v>94</v> (shared-string index 94)."""
    from openpyxl import Workbook, load_workbook
    from app.scrubber import scrub_xlsx
    wb = Workbook()
    ws = wb.active
    for i in range(120):
        ws.append([f"student note {i}", 94 if i == 5 else i])
    src = tmp_path / "g.xlsx"
    wb.save(src)
    out = tmp_path / "o.xlsx"
    scrub_xlsx(src, {"94": "[GRADE_5C94AB01DE2F]"}, out)
    rows = [[c.value for c in r] for r in load_workbook(out).active.iter_rows()]  # loads
    g = "[GRADE_5C94AB01DE2F]"
    assert [r[0] for r in rows] == [
        f"student note {g if i == 94 else i}" for i in range(120)]   # whole-number 94 only
    assert [r[1] for r in rows] == [g if i in (5, 94) else i for i in range(120)]


def test_sheet_named_after_a_family_round_trips(tmp_path):
    """Sheet titles can't hold [ or ]; the old code produced a workbook that
    would not open. Titles get the bracket-free token, references follow, and
    restoring puts the real name back everywhere."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    from app.mapper import EntityRegistry
    from app.scrubber import scrub_xlsx
    from app.unanonymize import restore_file
    from app.verifier import verify_output
    wb = Workbook()
    fam = wb.active
    fam.title = "Smith Family"
    fam.append(["Jane Smith", 500])
    summary = wb.create_sheet("Summary")
    summary["A1"] = "='Smith Family'!B1*2"
    wb.defined_names["smith_total"] = DefinedName("smith_total", attr_text="'Smith Family'!$B$1")
    src = tmp_path / "families.xlsx"
    wb.save(src)
    reg = EntityRegistry()
    reg.add("Smith Family", "ORG")
    reg.add("Jane Smith", "PERSON")
    out = tmp_path / "families_anon_ab12cd34.xlsx"
    result = scrub_xlsx(src, reg.as_replacement_map(), out)

    token = reg.replacements["Smith Family"].strip("[]")
    wb2 = load_workbook(out)                      # used to raise ValueError
    assert wb2.sheetnames == [token, "Summary"]
    assert wb2["Summary"]["A1"].value == f"='{token}'!B1*2"
    assert wb2.defined_names["smith_total"].attr_text == f"'{token}'!$B$1"
    assert verify_output(out, reg.as_replacement_map()).passed

    assert result.sheet_titles == {token: "Smith Family"}

    from app.restorer import build_index
    rmap = reg.as_replacement_map()
    for key in ({"replacement_map": rmap, "sheet_titles": result.sheet_titles},  # normal
                {"replacement_map": rmap}):                                    # token fallback
        res = restore_file(out, build_index([("k", key)]), display_name=out.name)
        wb3 = load_workbook(res.output_path)
        assert wb3.sheetnames == ["Smith Family", "Summary"]
        assert wb3["Smith Family"]["A1"].value == "Jane Smith"
        assert wb3["Summary"]["A1"].value == "='Smith Family'!B1*2"
        assert wb3.defined_names["smith_total"].attr_text == "'Smith Family'!$B$1"
        assert res.report.residual == 0


def test_ids_never_look_like_scientific_notation():
    from app.ids import is_valid_id
    assert not is_valid_id("123456789E12")
    assert not is_valid_id("1E2345678901")
    assert is_valid_id("12345678E12A")



def _formula_parses(formula: str) -> bool:
    from openpyxl.formula import Tokenizer
    try:
        Tokenizer(formula)
        return True
    except Exception:
        return False


def _anon_and_restore_titles(tmp_path, titles, rmap_values):
    """Build a workbook with `titles`, a Summary sheet referencing each, run
    anonymize -> restore through the real key file. Returns both workbooks."""
    import json as _json
    from openpyxl import Workbook, load_workbook
    from app.key_files import save_key_file
    from app.mapper import EntityRegistry
    from app.restorer import build_index
    from app.scrubber import scrub_xlsx
    from app.unanonymize import restore_file
    wb = Workbook()
    wb.remove(wb.active)
    for t in titles:
        wb.create_sheet(t)["B1"] = 7
    summ = wb.create_sheet("Summary")
    for i, t in enumerate(titles, start=1):
        summ.cell(row=i, column=1, value="='" + t.replace("'", "''") + "'!B1*2")
    src = tmp_path / "t.xlsx"
    wb.save(src)
    reg = EntityRegistry()
    for v, tag in rmap_values:
        reg.add(v, tag)
    out = tmp_path / "t_anon_ab12cd34.xlsx"
    result = scrub_xlsx(src, reg.as_replacement_map(), out)
    key_path = save_key_file("ab12cd34", "t.xlsx", {}, ["PERSON"], reg,
                             sheet_titles=result.sheet_titles)
    key = _json.loads(key_path.read_text())
    res = restore_file(out, build_index([(key_path.name, key)]), display_name=out.name)
    return load_workbook(out), load_workbook(res.output_path), key, res


def test_long_and_duplicate_sheet_titles_restore_exactly(tmp_path):
    """Review finding 3: a title that won't fit after substitution, and two
    titles that collide, must both come back exactly."""
    from app import ids
    # "2024 Pledges - PERSON_<id>" is 34 chars (won't fit); "Smith's" and
    # "Smith_s" both sanitize to "PERSON_<id>_s" (collide).
    titles = ["2024 Pledges - Smith", "Smith's", "Smith_s", "Smith Family"]
    anon, restored, key, res = _anon_and_restore_titles(
        tmp_path, titles, [("Smith", "PERSON"), ("Smith Family", "ORG")])
    for t in anon.sheetnames:
        assert "Smith" not in t and len(t) <= 31
    assert sum(t.startswith("SHEET_") for t in anon.sheetnames) == 2   # too long + collision
    assert restored.sheetnames == titles + ["Summary"]
    for i, t in enumerate(titles, start=1):
        f = restored["Summary"].cell(row=i, column=1).value
        assert f == "='" + t.replace("'", "''") + "'!B1*2" and _formula_parses(f)
    assert set(key["sheet_titles"].values()) == set(titles)
    # SHEET ids are released IDs too - never reissued.
    sheet_ids = {t[6:] for t in key["sheet_titles"] if t.startswith("SHEET_")}
    assert sheet_ids and sheet_ids <= ids.load_reserved()
    assert res.report.residual == 0


def test_apostrophe_sheet_title_restores_to_valid_formulas(tmp_path):
    """Review finding 4: 'O''Brien'!B1 must come back with the doubled quote."""
    anon, restored, _, _ = _anon_and_restore_titles(
        tmp_path, ["O'Brien"], [("O'Brien", "PERSON")])
    assert "O'Brien" not in anon.sheetnames
    assert restored.sheetnames == ["O'Brien", "Summary"]
    f = restored["Summary"]["A1"].value
    assert f == "='O''Brien'!B1*2" and _formula_parses(f)


def test_apostrophe_sheet_title_token_fallback_also_valid(tmp_path):
    """Same, when the key has no sheet_titles record (restore via the token)."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    from app.mapper import EntityRegistry
    from app.restorer import build_index
    from app.scrubber import scrub_xlsx
    from app.unanonymize import restore_file
    wb = Workbook()
    wb.active.title = "O'Brien"
    wb.active["B1"] = 7
    wb.create_sheet("Summary")["A1"] = "='O''Brien'!B1*2"
    wb.defined_names["ob"] = DefinedName("ob", attr_text="'O''Brien'!$B$1")
    src = tmp_path / "s.xlsx"
    wb.save(src)
    reg = EntityRegistry()
    reg.add("O'Brien", "PERSON")
    out = tmp_path / "o.xlsx"
    scrub_xlsx(src, reg.as_replacement_map(), out)
    res = restore_file(out, build_index([("k", {"replacement_map": reg.as_replacement_map()})]))
    wb2 = load_workbook(res.output_path)
    assert wb2.sheetnames == ["O'Brien", "Summary"]
    assert wb2["Summary"]["A1"].value == "='O''Brien'!B1*2"
    assert wb2.defined_names["ob"].attr_text == "'O''Brien'!$B$1"


def test_defined_names_keep_their_references(tmp_path):
    """A short texty value (B7) must not rewrite a defined name's reference."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    from app.scrubber import scrub_xlsx
    wb = Workbook()
    wb.active["B7"] = "B7"
    wb.defined_names["locker"] = DefinedName("locker", attr_text="Sheet!$B$7")
    src = tmp_path / "d.xlsx"
    wb.save(src)
    scrub_xlsx(src, {"B7": "[ID_B7B7B7B7B7B7]"}, tmp_path / "o.xlsx")
    wb2 = load_workbook(tmp_path / "o.xlsx")
    assert wb2.active["B7"].value == "[ID_B7B7B7B7B7B7]"
    assert wb2.defined_names["locker"].attr_text == "Sheet!$B$7"


def test_sheet_title_only_pii_is_detected_and_verified(tmp_path):
    """Review finding 7: a name that lives only in the sheet title reaches the
    model (extracted text) and the verifier."""
    from openpyxl import Workbook
    from app.extractors import extract
    from app.mapper import EntityRegistry
    from app.scrubber import scrub_xlsx
    from app.verifier import verify_output
    wb = Workbook()
    wb.active.title = "Maria Gonzalez IEP"
    wb.active.append(["goal", "status"])
    src = tmp_path / "iep.xlsx"
    wb.save(src)
    assert "Sheet: Maria Gonzalez IEP" in extract(src).text
    reg = EntityRegistry()
    reg.add("Maria Gonzalez", "PERSON")
    out = tmp_path / "o.xlsx"
    scrub_xlsx(src, reg.as_replacement_map(), out)
    assert "Maria" not in extract(out).text
    assert verify_output(out, reg.as_replacement_map()).passed
    # And a title that was NOT scrubbed fails verification.
    assert not verify_output(src, reg.as_replacement_map()).passed


def test_default_single_sheet_adds_no_header_line(tmp_path):
    from openpyxl import Workbook
    from app.extractors import extract
    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.save(tmp_path / "d.xlsx")
    assert extract(tmp_path / "d.xlsx").text.startswith("a\tb")


def test_markdown_and_snake_case_forms_restore():
    """Review finding 5."""
    for token, expected in [
        (f"_PERSON_{JANE}_", "_Jane Smith_"),
        (f"__PERSON_{JANE}__", "__Jane Smith__"),
        (f"PERSON_{JANE}_email", "Jane Smith_email"),
        (f"PERSON_{JANE}s", "Jane Smiths"),
        (f"*[PERSON_{JANE}]*", "*Jane Smith*"),
    ]:
        out, rep = _restore(token)
        assert out == expected, token
        assert rep.restored == 1 and rep.unresolved_count == 0


def test_known_id_that_cannot_be_restored_is_reported_not_silent():
    out, rep = _restore(f"ref x{JANE}9 and [PERSON_{BOB}]")
    assert out == f"ref x{JANE}9 and Bob Torres"
    assert rep.restored == 1
    assert rep.unresolved == [JANE]


def test_formula_reference_equal_to_a_pii_value_does_not_block_release(tmp_path):
    """Web-agent finding: a PII value "B7" also used as a cell reference in
    "Jane Smith"&B7 made verification fail forever on a correct file."""
    from openpyxl import Workbook, load_workbook
    from app.scrubber import scrub_xlsx
    from app.verifier import verify_output
    wb = Workbook()
    ws = wb.active
    ws["B7"] = "B7"
    ws["A1"] = '="Jane Smith"&B7'
    src = tmp_path / "f.xlsx"
    wb.save(src)
    rmap = {"Jane Smith": "[PERSON_3A4F9C2B1D0E]", "B7": "[ID_B7B7B7B7B7B7]"}
    out = tmp_path / "o.xlsx"
    scrub_xlsx(src, rmap, out)
    ws2 = load_workbook(out).active
    assert ws2["A1"].value == '="[PERSON_3A4F9C2B1D0E]"&B7'
    assert verify_output(out, rmap).passed
    # A literal left in a formula is still caught.
    leaky = tmp_path / "leak.xlsx"
    wb2 = Workbook()
    wb2.active["A1"] = '="Jane Smith"&B7'
    wb2.save(leaky)
    assert not verify_output(leaky, rmap).passed


def test_overlapping_detections_leave_no_fragment(tmp_path):
    """Web-agent finding: "Patient Jane" + "Jane Smith" in "Patient Jane
    Smith" left "Smith" behind whichever won. The union becomes one value."""
    import json as _json
    from unittest.mock import patch as _patch
    from app import detector, pipeline
    from app.config import UPLOADS_DIR
    from app.restorer import build_index, restore_text

    up = UPLOADS_DIR / "note.txt"
    up.write_text("Patient Jane Smith visited. Jane Smith called. Patient Jane waited.")
    sess = pipeline.new_session(up, "note.txt", ["PERSON"],
                                endpoint={"base_url": "http://localhost:1", "model": "m",
                                          "api_style": "ollama", "nickname": "t"},
                                custom_terms=[("PERSON", "Jane Smith")])
    fake = _json.dumps([{"text": "Patient Jane", "type": "PERSON"}])
    with _patch.object(detector.llm, "llm_call", return_value=fake):
        pipeline.run_extract_and_detect(sess)
    assert "Patient Jane Smith" in sess.registry.replacements
    pipeline.confirm_and_scrub(sess)
    assert sess.verify_result["passed"]
    out = sess.output_path.read_text()
    assert "Smith" not in out and "Jane" not in out
    idx = build_index([("k", {"replacement_map": sess.registry.as_replacement_map()})])
    assert restore_text(out, idx)[0] == "Patient Jane Smith visited. Jane Smith called. Patient Jane waited."


def test_overlap_unions_ignores_nesting_and_cross_cell_stretches():
    from app.replacer import overlap_unions
    assert overlap_unions("Jane Smith", ["Jane", "Jane Smith", "Smith"]) == []
    assert overlap_unions("Ann Lee\tLee Park", ["Ann Lee", "Lee\tLee Park"]) == []
    assert overlap_unions("Mary Ann Lee", ["Mary Ann", "Ann Lee"]) == [("Mary Ann Lee", ["Mary Ann", "Ann Lee"])]


def test_sheet_titles_that_differ_only_by_case_stay_distinct(tmp_path):
    """Excel sheet names are unique ignoring case: PERSON_X_s vs PERSON_X_S."""
    titles = ["Smith's", "Smith_S"]
    anon, restored, _, res = _anon_and_restore_titles(tmp_path, titles, [("Smith", "PERSON")])
    assert len({t.casefold() for t in anon.sheetnames}) == len(anon.sheetnames)
    assert restored.sheetnames == titles + ["Summary"]
    assert res.report.residual == 0
