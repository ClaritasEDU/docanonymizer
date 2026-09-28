"""Community identifiers on rosters, end to end through the Flask API.

A stub Family Graph on a loopback port serves responses captured from the
real Family Graph roster engine (tests/fixtures/fg_*.json, produced by its
roster code on the same sheet), so these tests hold docanonymizer to the
actual wire contract.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

import pytest

FIX = Path(__file__).parent / "fixtures"
SHEET = json.loads((FIX / "fg_request_sheet.json").read_text())
PLAN = json.loads((FIX / "fg_plan.json").read_text())
COMMIT = json.loads((FIX / "fg_commit.json").read_text())
REFUSED = json.loads((FIX / "fg_commit_refused.json").read_text())
DECISION = json.loads((FIX / "fg_decision.json").read_text())
API_KEY = "sk_test_roster_key_0123456789abcdef"

NAMES = ["Emma", "Liam", "Jane", "Ava", "Marie"]


# ---------------------------------------------------------------------------
# Stub Family Graph
# ---------------------------------------------------------------------------

class StubFG:
    def __init__(self):
        self.requests: list[dict] = []
        self.mode = "normal"            # normal | always_refuse | forbidden
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), self._handler())
        self.port = self.server.server_address[1]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    @property
    def url(self):
        return f"http://127.0.0.1:{self.port}"

    def close(self):
        self.server.shutdown()
        self.server.server_close()

    def _handler(self):
        stub = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):  # quiet
                pass

            def _send(self, status, body):
                data = json.dumps(body).encode()
                self.send_response(status)
                self.send_header("content-type", "application/json")
                self.send_header("content-length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def _authed(self):
                return self.headers.get("authorization") == f"Bearer {API_KEY}"

            def do_GET(self):
                if self.path == "/api/health":
                    return self._send(200, {"status": "ok"})
                if self.path.startswith("/api/identity/roster/lookup/"):
                    return self._send(404 if self._authed() else 401, {"error": "x"})
                self._send(404, {})

            def do_POST(self):
                length = int(self.headers.get("content-length") or 0)
                body = json.loads(self.rfile.read(length) or b"{}")
                stub.requests.append({"path": self.path, "body": body, "headers": dict(self.headers)})
                if not self._authed():
                    return self._send(401, {"error": "unauthorized"})
                if stub.mode == "forbidden":
                    return self._send(403, {"error": "forbidden"})
                if self.path == "/api/identity/roster/plan":
                    return self._send(200, PLAN)
                if self.path == "/api/identity/roster/commit":
                    d = (body.get("decisions") or {}).get(DECISION["key"])
                    if stub.mode == "always_refuse" or not d or d.get("target") != DECISION["target"]:
                        return self._send(409, REFUSED)
                    return self._send(201, COMMIT)
                self._send(404, {})

        return H


@pytest.fixture
def fg():
    stub = StubFG()
    yield stub
    stub.close()


def _client():
    from app.server import create_app
    app = create_app()
    app.config["TESTING"] = True
    return app.test_client()


def _connect(c, url):
    r = c.post("/api/familygraph", json={"base_url": url, "api_key": API_KEY, "category": "school"})
    assert r.status_code == 200, r.get_json()
    return r


def _roster_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(SHEET["headers"])
    w.writerows(SHEET["rows"])
    return buf.getvalue().encode()


# The model finds most names but misses some: Liam, Ava and Marie are never
# flagged. The community layer must still cover their cells.
LLM = json.dumps([
    {"text": "Emma", "type": "PERSON"}, {"text": "Smith", "type": "PERSON"},
    {"text": "Jane", "type": "PERSON"}, {"text": "Garcia", "type": "PERSON"},
    {"text": "jane@example.org", "type": "EMAIL"}, {"text": "Ann", "type": "PERSON"},
])


def _detect(c, data: bytes, name: str):
    from app import detector
    with patch.object(detector.llm, "llm_call", return_value=LLM):
        r = c.post("/api/anonymize/upload", data={
            "file": (io.BytesIO(data), name), "tags": "ALL",
        }, content_type="multipart/form-data")
        assert r.status_code == 200, r.get_json()
        sid = r.get_json()["session_id"]
        for _ in range(200):
            s = c.get(f"/api/anonymize/{sid}/status").get_json()
            if s.get("detection_complete") or s.get("error"):
                break
            time.sleep(0.05)
    assert s.get("detection_complete"), s
    return sid, s


def _wait_results(c, sid):
    for _ in range(300):
        r = c.get(f"/api/anonymize/{sid}/results").get_json()
        if r.get("verify_result") or r.get("error"):
            return r
        time.sleep(0.05)
    raise AssertionError("scrub did not finish")


def _run_roster(c, data, name):
    sid, _ = _detect(c, data, name)
    view = c.get(f"/api/anonymize/{sid}/community").get_json()
    assert view["status"] == "ready" and view["pending"] == 1
    r = c.post(f"/api/anonymize/{sid}/community/decide",
               json={"key": DECISION["key"], "action": "attach", "target": DECISION["target"]})
    assert r.status_code == 200 and r.get_json()["pending"] == 0
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
    assert r.status_code == 200, r.get_json()
    return sid, _wait_results(c, sid)


def _ids(result, given):
    for row in result["sheets"][0]["rows"]:
        for p in row.get("persons") or []:
            if p["given_name"] == given:
                return p["community_id"]
    raise KeyError(given)


# ---------------------------------------------------------------------------
# Connection settings
# ---------------------------------------------------------------------------

def test_settings_never_return_the_key_and_file_is_private(fg):
    from app.config import FAMILYGRAPH_FILE
    c = _client()
    _connect(c, fg.url)
    got = c.get("/api/familygraph").get_json()
    assert got == {"configured": True, "base_url": fg.url, "key_set": True, "category": "school"}
    assert API_KEY not in json.dumps(got)
    assert (FAMILYGRAPH_FILE.stat().st_mode & 0o777) == 0o600
    # Saving again without a key keeps the stored one.
    assert c.post("/api/familygraph", json={"base_url": fg.url, "category": "church"}).status_code == 200
    assert json.loads(FAMILYGRAPH_FILE.read_text())["api_key"] == API_KEY
    assert c.post("/api/familygraph/test").get_json() == {"status": "ok"}
    assert c.delete("/api/familygraph").status_code == 204
    assert c.get("/api/familygraph").get_json()["configured"] is False


def test_a_non_local_family_graph_is_refused_with_no_override():
    c = _client()
    for url in ("http://8.8.8.8:3500", "https://example.com", "http://127.0.0.1:3500/api"):
        r = c.post("/api/familygraph", json={"base_url": url, "api_key": API_KEY, "allow_nonlocal": True})
        assert r.status_code == 400, url


def test_familygraph_json_is_gitignored():
    root = Path(__file__).resolve().parent.parent
    assert "familygraph.json" in (root / ".gitignore").read_text().split()


def test_wrong_key_is_reported_plainly(fg):
    c = _client()
    c.post("/api/familygraph", json={"base_url": fg.url, "api_key": "sk_wrong_key_0123456789abc"})
    assert c.post("/api/familygraph/test").get_json() == {"status": "err", "error": "Family Graph rejected the API key"}


# ---------------------------------------------------------------------------
# The roster flow
# ---------------------------------------------------------------------------

def test_csv_roster_end_to_end(fg):
    from app.config import KEYS_DIR, OUTPUT_DIR
    c = _client()
    _connect(c, fg.url)
    sid, _ = _detect(c, _roster_csv(), "roster.csv")

    # Confirm is refused while a person still needs a decision.
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
    assert r.status_code == 409 and "decision" in r.get_json()["error"]
    view = c.get(f"/api/anonymize/{sid}/community").get_json()
    review = [i for i in view["items"] if i["action"] == "review"]
    assert [i["key"] for i in review] == [DECISION["key"]]
    assert review[0]["label"] == "Marie Garcia"
    assert review[0]["candidates"][0]["community_id"] == DECISION["target"]

    # A target that is not a listed candidate is refused.
    bad = c.post(f"/api/anonymize/{sid}/community/decide",
                 json={"key": DECISION["key"], "action": "attach", "target": "I0000000000000000"})
    assert bad.status_code == 400
    c.post(f"/api/anonymize/{sid}/community/decide",
           json={"key": DECISION["key"], "action": "attach", "target": DECISION["target"]})
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
    assert r.status_code == 200
    res = _wait_results(c, sid)
    assert res["verify_result"]["passed"], res
    assert res["community"]["status"] == "committed"

    # What reached Family Graph: the rows, the decision, our key and actor.
    plan_req = next(x for x in fg.requests if x["path"].endswith("/plan"))
    commit_req = [x for x in fg.requests if x["path"].endswith("/commit")][-1]
    assert plan_req["headers"]["x-family-graph-actor"] == "docanonymizer"
    assert plan_req["body"]["sheets"][0]["rows"] == SHEET["rows"]
    assert plan_req["body"]["source_ref"] == f"docanonymizer:{sid}"
    assert commit_req["body"]["decisions"] == {DECISION["key"]: {"action": "attach", "target": DECISION["target"]}}

    out = OUTPUT_DIR / res["output_filename"]
    rows = list(csv.reader(out.read_text(encoding="utf-8").splitlines()))
    assert rows[0][-1] == "FAMILY_ID", "no household column on this sheet - one is appended"
    emma, jane = _ids(COMMIT, "Emma"), _ids(COMMIT, "Jane")
    assert rows[1][0] == f"[{emma}]"
    assert rows[1][3] == f"[{jane}]"
    assert rows[2][3] == f"[{jane}]", "one person, one id, every row"
    assert rows[3][3] == f"[{DECISION['target']}]", "the operator's decision is the id used"
    fam = COMMIT["sheets"][0]["rows"][0]["family"]["community_id"]
    assert rows[1][-1] == rows[2][-1] == f"[{fam}]"
    assert rows[3][-1] != rows[1][-1]
    # Names the model missed are still gone; surnames carry per-value tokens.
    body = out.read_text(encoding="utf-8")
    for name in NAMES + ["Smith", "Garcia"]:
        assert not re.search(rf"\b{name}\b", body), name
    assert re.fullmatch(r"\[PERSON_[0-9A-F]{12}\]", rows[1][1])

    key = json.loads((KEYS_DIR / res["key_filename"]).read_text())
    assert key["identity_registry"][jane] == {"kind": "person", "display": "Jane Smith", "family": fam}
    assert key["identity_registry"][fam]["kind"] == "family"
    assert {"sheet": 0, "row": 1, "col": 0, "written": f"[{emma}]", "original": "Emma"} in key["identity_cells"]
    assert key["identity_columns"] == [{"sheet": 0, "col": 7, "header_row": 0, "header": "FAMILY_ID"}]
    assert key["community_source"]["import_runs"] == COMMIT["import_runs"]
    # Community ids never enter this app's issued-id ledger.
    from app import ids
    assert not any(jane.strip("I") in i for i in ids.load_reserved())

    # Restoring the anonymized file gives back the original exactly.
    r = c.post("/api/unanonymize", data={
        "file": (io.BytesIO(out.read_bytes()), out.name),
        "keys": json.dumps([res["key_filename"]]),
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    restored = (OUTPUT_DIR / r.get_json()["output_filename"]).read_text(encoding="utf-8")
    assert list(csv.reader(restored.splitlines())) == [SHEET["headers"]] + SHEET["rows"]


def test_xlsx_roster_end_to_end(fg):
    from openpyxl import Workbook, load_workbook
    from app.config import KEYS_DIR, OUTPUT_DIR
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(SHEET["headers"])
    for row in SHEET["rows"]:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    c = _client()
    _connect(c, fg.url)
    sid, res = _run_roster(c, buf.getvalue(), "roster.xlsx")
    assert res["verify_result"]["passed"], res
    out = OUTPUT_DIR / res["output_filename"]
    ws2 = load_workbook(out).active
    assert ws2.cell(row=1, column=8).value == "FAMILY_ID"
    assert ws2.cell(row=2, column=1).value == f"[{_ids(COMMIT, 'Emma')}]"
    assert ws2.cell(row=4, column=4).value == f"[{DECISION['target']}]"

    r = c.post("/api/unanonymize", data={
        "file": (io.BytesIO(out.read_bytes()), out.name),
        "keys": json.dumps([res["key_filename"]]),
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    ws3 = load_workbook(OUTPUT_DIR / r.get_json()["output_filename"]).active
    got = [[("" if v is None else str(v)) for v in row] for row in ws3.iter_rows(values_only=True)]
    assert got == [SHEET["headers"]] + SHEET["rows"]
    assert ws3.max_column == len(SHEET["headers"]), "the appended column is gone"


def test_ai_answer_with_community_ids_restores_to_names(fg):
    c = _client()
    _connect(c, fg.url)
    _, res = _run_roster(c, _roster_csv(), "roster.csv")
    jane, emma = _ids(COMMIT, "Jane"), _ids(COMMIT, "Emma")
    fam = COMMIT["sheets"][0]["rows"][0]["family"]["community_id"]
    answer = (f"Top volunteers: [{jane}], {emma.lower()}, \\[{jane}\\], PARENT_{jane}, "
              f"and household [{fam}] (bare hex {fam[1:]}). Unknown [I0123456789ABCDEF].")
    r = c.post("/api/unanonymize/text", json={"text": answer, "keys": [res["key_filename"]]})
    body = r.get_json()
    assert body["text"] == ("Top volunteers: Jane Smith, Emma Smith, Jane Smith, Jane Smith, "
                            "and household Smith Family (bare hex Smith Family). Unknown [I0123456789ABCDEF].")
    assert body["report"]["by_tag"]["INDIVIDUAL"] == 4
    assert body["report"]["by_tag"]["FAMILY"] == 2
    assert body["report"]["unresolved"] == ["[I0123456789ABCDEF]"]


def test_commit_that_finds_new_items_returns_to_review(fg):
    c = _client()
    _connect(c, fg.url)
    sid, _ = _detect(c, _roster_csv(), "roster.csv")
    c.post(f"/api/anonymize/{sid}/community/decide",
           json={"key": DECISION["key"], "action": "attach", "target": DECISION["target"]})
    fg.mode = "always_refuse"
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
    assert r.status_code == 409
    assert "new items" in r.get_json()["error"]
    assert c.get(f"/api/anonymize/{sid}/results").get_json()["scrub_steps"] == [], "nothing was scrubbed"


def test_family_graph_down_blocks_until_the_operator_chooses(fg):
    from app.config import OUTPUT_DIR
    c = _client()
    _connect(c, fg.url)
    fg.close()                                   # Family Graph goes away
    sid, _ = _detect(c, _roster_csv(), "roster.csv")
    view = c.get(f"/api/anonymize/{sid}/community").get_json()
    assert view["status"] == "error" and "not reachable" in view["error"]
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []})
    assert r.status_code == 409 and "Family Graph unavailable" in r.get_json()["error"]
    # Explicitly continuing without community ids: the normal pipeline runs.
    assert c.post(f"/api/anonymize/{sid}/community/skip", json={"skip": True}).status_code == 200
    assert c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": []}).status_code == 200
    res = _wait_results(c, sid)
    assert res["verify_result"]["passed"]
    assert res["community"]["status"] == "skipped"
    assert "[I" not in (OUTPUT_DIR / res["output_filename"]).read_text(encoding="utf-8")


def test_scope_error_is_explained(fg):
    c = _client()
    _connect(c, fg.url)
    fg.mode = "forbidden"
    sid, _ = _detect(c, _roster_csv(), "roster.csv")
    view = c.get(f"/api/anonymize/{sid}/community").get_json()
    assert view["status"] == "error" and "scope" in view["error"]


def test_keeping_names_as_original_skips_community_ids(fg):
    from app.config import OUTPUT_DIR
    c = _client()
    _connect(c, fg.url)
    sid, _ = _detect(c, _roster_csv(), "roster.csv")
    r = c.post(f"/api/anonymize/{sid}/confirm", json={"deselected": [], "deselected_types": ["PERSON"]})
    assert r.status_code == 200
    res = _wait_results(c, sid)
    assert res["community"]["status"] == "skipped"
    assert not any(x["path"].endswith("/commit") for x in fg.requests), "nothing was written to Family Graph"
    assert "Emma" in (OUTPUT_DIR / res["output_filename"]).read_text(encoding="utf-8")


def test_documents_and_unconfigured_runs_are_untouched(fg):
    c = _client()
    sid, _ = _detect(c, _roster_csv(), "roster.csv")        # not connected
    assert c.get(f"/api/anonymize/{sid}/community").get_json()["status"] == "off"
    _connect(c, fg.url)
    sid, _ = _detect(c, b"Jane Smith wrote to Emma.", "memo.txt")
    assert c.get(f"/api/anonymize/{sid}/community").get_json()["status"] == "off"
    assert fg.requests == []


def test_a_name_the_model_missed_in_free_text_is_still_caught(fg):
    """Liam's first name also appears in the notes column. The model never
    flagged it; the community layer registers it before the preview."""
    c = _client()
    _connect(c, fg.url)
    sheet = json.loads(json.dumps(SHEET))
    sheet["rows"][1][6] = "Liam rides the bus"
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(sheet["headers"])
    w.writerows(sheet["rows"])
    sid, status = _detect(c, buf.getvalue().encode(), "roster.csv")
    assert any(s["original"] == "Liam" for s in status["preview"]["spans"])


# ---------------------------------------------------------------------------
# Engine-level guarantees
# ---------------------------------------------------------------------------

def test_short_values_are_never_written_inside_a_community_token():
    from app.replacer import LiteralReplacer, apply
    token = "[IA4F9C2B1D0E7F21A]"
    r = LiteralReplacer({"A4": "[SID_111111111111]", "94": "[GRADE_222222222222]", "F9C2": "[ID_333333333333]"},
                        extra_protected={token, token.strip("[]")})
    assert apply(f"{token} A4 94", r) == f"{token} [SID_111111111111] [GRADE_222222222222]"


def test_restorer_community_forms_and_variants():
    from app.restorer import build_index, restore_text
    k1 = {"session_id": "aaaa1111", "original_filename": "a.xlsx", "replacement_map": {}, "created_at": "2026-01-01T00:00:00Z",
          "identity_registry": {"I00000000000000AA": {"kind": "person", "display": "Robert Jones"}}}
    k2 = {"session_id": "bbbb2222", "original_filename": "b.xlsx", "replacement_map": {}, "created_at": "2026-09-01T00:00:00Z",
          "identity_registry": {"I00000000000000AA": {"kind": "person", "display": "Bob Jones"}}}
    idx = build_index([("a", k1), ("b", k2)])
    assert idx.community_variants == 1
    text, rep = restore_text("[I00000000000000AA] / [f00000000000000aa] / 00000000000000AA", idx)
    assert text == "Bob Jones / Bob Jones / Bob Jones", "newest spelling wins; hex alone restores"
    assert rep.relabeled == 1 and rep.untagged == 1
    # Order of keys does not change the winner.
    idx2 = build_index([("b", k2), ("a", k1)])
    assert restore_text("[I00000000000000AA]", idx2)[0] == "Bob Jones"


def test_log_holds_no_names_emails_or_key(fg):
    from app.config import LOG_FILE
    c = _client()
    _connect(c, fg.url)
    _run_roster(c, _roster_csv(), "roster.csv")
    log = LOG_FILE.read_text(encoding="utf-8")
    for s in NAMES + ["Smith", "Garcia", "jane@example.org", API_KEY, DECISION["target"]]:
        assert s not in log, s
    assert "community commit" in log
