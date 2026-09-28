"""Regression tests for the community-id core fixes (2026-09-28 review).

Each test drives the real pipeline (extract -> overrides -> scrub -> verify
-> key file -> restore) with a hand-built Family Graph commit result, so no
Family Graph server is needed. Names here are fixtures, never real people.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile

import pytest

ANN, BOB, EMMA, LIAM = "I1111222233334444", "I5555666677778888", "I4AD4825DB3958878", "IE73AAC749C2913BA"
SMITHS = "F3D6B57AFBDADAEBD"
NOBODY = "F9B0C11D2E3F4A5B6"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _person(key, given, family, cells, cid, action="new"):
    return {"key": key, "slot": int(key.split(":")[2]), "given_name": given, "family_name": family,
            "name_cells": cells, "action": action, "community_id": cid}


def _session(name: str, data: bytes, result: dict, flagged=(), sid="ab12cd34"):
    """A detected session whose roster Family Graph has committed."""
    from app import community, pipeline
    from app.config import UPLOADS_DIR
    from app.extractors import extract
    from app.mapper import EntityRegistry
    up = UPLOADS_DIR / name
    up.write_bytes(data)
    ex = extract(up)
    reg = EntityRegistry()
    for v in flagged:
        reg.add(v, "PERSON")
    sess = pipeline.Session(id=sid, upload_path=up, original_filename=name, endpoint={},
                            extract=ex, registry=reg, detection_complete=True)
    st = community.CommunityState(status="committed", result=result, plan=result)
    st.sheets = community.build_sheets(community.tables_of(ex))
    sess.community = st
    return sess


def _scrub(sess):
    from app import pipeline
    pipeline.confirm_and_scrub(sess)
    assert sess.error is None, sess.error
    assert sess.verify_result and sess.verify_result["passed"], sess.verify_result
    return sess.output_path, json.loads(sess.key_path.read_text())


def _restore(out_path, key: dict, name: str):
    """Restore the anonymized file uploaded under `name` (as the server does)."""
    from app import unanonymize
    from app.config import UPLOADS_DIR
    from app.restorer import build_index
    up = UPLOADS_DIR / name
    up.write_bytes(out_path.read_bytes())
    res = unanonymize.restore_file(up, build_index([("key", key)]), display_name=name)
    return res


def _csv_bytes(rows) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)            # the csv module's own \r\n
    return buf.getvalue().encode()


def _grid(path):
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    return [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# 11. A merged household cell must not crash the scrub after the commit
# ---------------------------------------------------------------------------

def _merged_roster() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Family", "Child"])
    ws.append(["Smith Family", "Ann Smith"])
    ws.append([None, "Bob Smith"])
    ws.merge_cells("A2:A3")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _merged_result():
    fam = lambda i: {"key": f"0:{i}:family", "display_name": "Smith Family", "cell": {"col": 0},
                     "action": "new" if i == 0 else "matched", "community_id": SMITHS}
    return {"committed": True, "import_runs": ["imp_merged000001"], "summary": {},
            "sheets": [{"rows": [
                {"index": 0, "persons": [_person("0:0:0", "Ann", "Smith", [{"col": 1, "part": "full"}], ANN)],
                 "family": fam(0)},
                {"index": 1, "persons": [_person("0:1:0", "Bob", "Smith", [{"col": 1, "part": "full"}], BOB)],
                 "family": fam(1)}]}]}


def test_merged_household_column_scrubs_and_restores_exactly():
    sess = _session("merged.xlsx", _merged_roster(), _merged_result(),
                    flagged=("Ann Smith", "Bob Smith", "Smith Family"))
    out, key = _scrub(sess)
    got = _grid(out)
    assert got[0] == ["Family", "Child", "FAMILY_ID"]
    assert got[1][:2] == [f"[{SMITHS}]", f"[{ANN}]"]
    # The empty (merged) cell never takes a write; its row's household id
    # lives in the appended column instead, so nothing is lost.
    assert got[2] == ["", f"[{BOB}]", f"[{SMITHS}]"]
    from openpyxl import load_workbook
    assert "A2:A3" in {str(r) for r in load_workbook(out).active.merged_cells.ranges}
    assert not any(c["original"] == "" for c in key["identity_cells"])

    res = _restore(out, key, out.name)
    assert _grid(res.output_path) == [["Family", "Child"], ["Smith Family", "Ann Smith"], ["", "Bob Smith"]]


def test_scrub_xlsx_skips_a_merged_cell_override_instead_of_crashing(tmp_path):
    from app.scrubber import scrub_xlsx
    src = tmp_path / "in.xlsx"
    src.write_bytes(_merged_roster())
    overrides = {(0, 1, 0): ("Smith Family", f"[{SMITHS}]"), (0, 2, 0): ("", f"[{SMITHS}]")}
    result = scrub_xlsx(src, {}, tmp_path / "out.xlsx", cell_overrides=overrides)
    assert result.overrides_skipped == [(0, 2, 0)]
    assert _grid(tmp_path / "out.xlsx")[1][0] == f"[{SMITHS}]"


def test_a_scrub_failure_after_the_commit_says_the_ids_are_safe(monkeypatch):
    from app import pipeline
    sess = _session("merged.xlsx", _merged_roster(), _merged_result(), flagged=("Ann Smith",))

    def boom(*a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(pipeline, "scrub_xlsx", boom)
    pipeline.confirm_and_scrub(sess)
    assert sess.error.startswith("scrub failed: disk full")
    assert "Family Graph already saved" in sess.error and "same ids" in sess.error
    assert not sess.upload_path.exists(), "temp hygiene still holds"


# ---------------------------------------------------------------------------
# 16. Exact restore must not depend on the download filename
# ---------------------------------------------------------------------------

ROSTER = [["Student First", "Student Last", "Parent First", "Parent Last", "Notes"],
          ["Emma", "Smith", "Ann", "Smith", "Allergic to peanuts"],
          ["Liam", "Smith", "Ann", "Smith", ""]]


def _roster_result():
    fam = lambda i: {"key": f"0:{i}:family", "display_name": "Smith Family", "cell": None,
                     "action": "new", "community_id": SMITHS}
    gf = lambda g, f: [{"col": g, "part": "given"}, {"col": f, "part": "family"}]
    return {"committed": True, "import_runs": ["imp_roster000001"], "summary": {},
            "sheets": [{"rows": [
                {"index": 0, "persons": [_person("0:0:0", "Ann", "Smith", gf(2, 3), ANN),
                                         _person("0:0:1", "Emma", "Smith", gf(0, 1), EMMA)], "family": fam(0)},
                {"index": 1, "persons": [_person("0:1:0", "Ann", "Smith", gf(2, 3), ANN, "matched"),
                                         _person("0:1:1", "Liam", "Smith", gf(0, 1), LIAM)], "family": fam(1)}]}]}


@pytest.mark.parametrize("rename", ["{stem} (1){ext}", "{stem}_1{ext}", "renamed{ext}", "my roster{ext}"])
def test_exact_restore_survives_a_renamed_download(rename):
    from werkzeug.utils import secure_filename
    sess = _session("roster.csv", _csv_bytes(ROSTER), _roster_result(), flagged=("Smith",))
    out, key = _scrub(sess)
    name = secure_filename(rename.format(stem=out.stem, ext=out.suffix))
    res = _restore(out, key, name)
    assert list(csv.reader(res.output_path.read_text(encoding="utf-8").splitlines())) == ROSTER
    assert "FAMILY_ID" not in res.output_path.read_text(encoding="utf-8")
    assert res.output_path.name.startswith(("roster_restored", "renamed_restored", "my_roster_restored"))


def test_exact_restore_picks_the_key_whose_cells_are_in_the_file():
    """Two roster keys selected; the file came from the second one."""
    from app import unanonymize
    from app.config import UPLOADS_DIR
    from app.restorer import build_index
    other_rows = [ROSTER[0], ["Emma", "Smith", "Ann", "Smith", "Walks home"]]
    other = _roster_result()
    other["sheets"][0]["rows"] = other["sheets"][0]["rows"][:1]
    s1 = _session("other.csv", _csv_bytes(other_rows), other, flagged=("Smith",), sid="11112222")
    _, key1 = _scrub(s1)
    s2 = _session("roster.csv", _csv_bytes(ROSTER), _roster_result(), flagged=("Smith",), sid="33334444")
    out2, key2 = _scrub(s2)
    up = UPLOADS_DIR / "download.csv"
    up.write_bytes(out2.read_bytes())
    res = unanonymize.restore_file(up, build_index([("k1", key1), ("k2", key2)]), display_name="download.csv")
    assert list(csv.reader(res.output_path.read_text(encoding="utf-8").splitlines())) == ROSTER


# ---------------------------------------------------------------------------
# 17. A surname from a full-name or list cell is a guaranteed catch
# ---------------------------------------------------------------------------

def _full_name_result(cells_part="full"):
    return {"committed": True, "import_runs": ["imp_fullname0001"], "summary": {},
            "sheets": [{"rows": [
                {"index": 0, "persons": [_person("0:0:0", "Emma", "Smith", [{"col": 0, "part": cells_part}], EMMA)],
                 "family": {"key": "0:0:family", "display_name": "Smith Family", "cell": None,
                            "action": "new", "community_id": SMITHS}}]}]}


@pytest.mark.parametrize("part", ["full", "list"])
def test_a_surname_from_a_full_name_cell_is_caught_in_free_text(part):
    from app import community
    rows = [["Student", "Notes"], ["Emma Smith", "Smith family moving in May"]]
    # The model flagged the full name only.
    sess = _session("full.csv", _csv_bytes(rows), _full_name_result(part), flagged=("Emma Smith",))
    community._register_guaranteed_catches(sess, sess.community)
    assert "Smith" in sess.registry.replacements
    out, key = _scrub(sess)
    got = list(csv.reader(out.read_text(encoding="utf-8").splitlines()))
    assert got[1][0] == f"[{EMMA}]", "the name cell still carries the community id, whole"
    assert not re.search(r"\bSmith\b", out.read_text(encoding="utf-8"))
    assert re.fullmatch(r"\[PERSON_[0-9A-F]{12}\] family moving in May", got[1][1])
    res = _restore(out, key, out.name)
    assert list(csv.reader(res.output_path.read_text(encoding="utf-8").splitlines())) == rows


def test_a_surname_not_in_free_text_is_not_added():
    from app import community
    rows = [["Student", "Notes"], ["Emma Smith", "Smithson visited"]]
    sess = _session("full.csv", _csv_bytes(rows), _full_name_result(), flagged=())
    community._register_guaranteed_catches(sess, sess.community)
    assert "Smith" not in sess.registry.replacements, "whole words only"


# ---------------------------------------------------------------------------
# 18. A household with no display name is never restored to ""
# ---------------------------------------------------------------------------

def test_a_nameless_household_gets_a_label_and_empty_displays_are_unresolved():
    from app import community
    from app.extractors import ExtractResult
    from app.restorer import build_index, restore_text
    # A one-column sheet is found only under an exact person-name header.
    table = [["Name"], ["Ann"], ["Ben"]]
    fam = {"key": "0:0:family", "display_name": None, "cell": None, "action": "new", "community_id": NOBODY}
    st = community.CommunityState(status="committed")
    st.sheets = community.build_sheets([table])
    st.result = {"sheets": [{"rows": [
        {"index": 0, "persons": [_person("0:0:0", "Ann", None, [{"col": 0, "part": "given"}], ANN)], "family": fam},
        {"index": 1, "persons": [_person("0:1:0", "Ben", None, [{"col": 0, "part": "given"}], BOB)],
         "family": dict(fam, key="0:1:family")}]}]}
    ex = ExtractResult(text="", char_count=0, original_suffix="csv", payload={"rows": table})
    ov = community.overrides_for(type("S", (), {"community": st, "extract": ex})())
    assert ov.registry[NOBODY] == {"kind": "family", "display": "Ann & Ben household"}

    # A key written before this fix: the empty label is reported, never deleted.
    key = {"session_id": "aaaa0000", "replacement_map": {}, "identity_registry": {
        NOBODY: {"kind": "family", "display": ""}, ANN: {"kind": "person", "display": "Ann"}}}
    out, rep = restore_text(f"[{NOBODY}] owes $200; [{ANN}] too; bare {NOBODY[1:]}", build_index([("k", key)]))
    assert out == f"[{NOBODY}] owes $200; Ann too; bare {NOBODY[1:]}"
    assert rep.restored == 1 and rep.unresolved == [f"[{NOBODY}]", NOBODY[1:]]


def test_a_household_label_from_one_row_is_not_blanked_by_the_next():
    from app import community
    from app.extractors import ExtractResult
    table = [["First", "Last"], ["Ann", "Smith"], ["Ben", ""]]
    st = community.CommunityState(status="committed")
    st.sheets = community.build_sheets([table])
    fam = lambda i, d: {"key": f"0:{i}:family", "display_name": d, "cell": None, "action": "new",
                        "community_id": SMITHS}
    st.result = {"sheets": [{"rows": [
        {"index": 0, "persons": [_person("0:0:0", "Ann", "Smith", [{"col": 0, "part": "given"}], ANN)],
         "family": fam(0, "Smith Family")},
        {"index": 1, "persons": [_person("0:1:0", "Ben", None, [{"col": 0, "part": "given"}], BOB)],
         "family": fam(1, None)}]}]}
    ex = ExtractResult(text="", char_count=0, original_suffix="csv", payload={"rows": table})
    ov = community.overrides_for(type("S", (), {"community": st, "extract": ex})())
    assert ov.registry[SMITHS]["display"] == "Smith Family"


def test_an_empty_newer_spelling_never_replaces_a_real_one():
    from app.restorer import build_index
    old = {"session_id": "a", "created_at": "2026-01-01T00:00:00Z", "replacement_map": {},
           "identity_registry": {SMITHS: {"kind": "family", "display": "Smith Family"}}}
    new = {"session_id": "b", "created_at": "2026-02-01T00:00:00Z", "replacement_map": {},
           "identity_registry": {SMITHS: {"kind": "family", "display": ""}}}
    for order in ((old, new), (new, old)):
        idx = build_index([(str(i), k) for i, k in enumerate(order)])
        assert idx.community[SMITHS[1:]] == ("F", "Smith Family")


# ---------------------------------------------------------------------------
# 19. A tagged community id with the letter dropped still restores
# ---------------------------------------------------------------------------

def test_relabeled_community_ids_without_the_letter_restore():
    from app.restorer import build_index, restore_text
    key = {"session_id": "cccc1111", "replacement_map": {"Jane Doe": "[PERSON_3A4F9C2B1D0E]"},
           "identity_registry": {SMITHS: {"kind": "family", "display": "Smith Family"},
                                 EMMA: {"kind": "person", "display": "Emma Smith"}}}
    idx = build_index([("k", key)])
    fh, eh = SMITHS[1:], EMMA[1:]
    text = (f"FAMILY_{fh} / [PERSON_{eh}] / DONOR_{eh} / \\[STUDENT\\_{eh}\\] / "
            f"family_{fh.lower()} / [PERSON_3A4F9C2B1D0E]")
    out, rep = restore_text(text, idx)
    assert out == ("Smith Family / Emma Smith / Emma Smith / Emma Smith / Smith Family / Jane Doe")
    assert rep.restored == 6 and rep.relabeled == 5 and rep.unresolved == []


def test_an_unknown_or_glued_community_hex_is_reported_never_silent():
    from app.restorer import build_index, restore_text
    key = {"session_id": "cccc1111", "replacement_map": {},
           "identity_registry": {EMMA: {"kind": "person", "display": "Emma Smith"}}}
    idx = build_index([("k", key)])
    text = f"FAMILY_0123456789ABCDEF and DONOR_0123456789ABCDEF and x{EMMA[1:]}9"
    out, rep = restore_text(text, idx)
    assert out == text
    assert rep.unresolved == ["FAMILY_0123456789ABCDEF", "DONOR_0123456789ABCDEF", EMMA[1:]]


# ---------------------------------------------------------------------------
# 26. The appended FAMILY_ID column never pads short or blank CSV rows
# ---------------------------------------------------------------------------

def test_csv_round_trip_with_blank_and_short_rows_is_byte_exact():
    rows = [["Name", "Grade", "Room"], ["Emma", "3", "B"], [], ["Total", "3"]]
    data = _csv_bytes(rows)
    result = {"committed": True, "import_runs": ["imp_short0000001"], "summary": {},
              "sheets": [{"rows": [
                  {"index": 0, "persons": [_person("0:0:0", "Emma", None, [{"col": 0, "part": "given"}], EMMA)],
                   "family": {"key": "0:0:family", "display_name": "Emma household", "cell": None,
                              "action": "new", "community_id": SMITHS}}]}]}
    sess = _session("short.csv", data, result, flagged=("Emma",))
    out, key = _scrub(sess)
    got = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8"))))
    assert got == [["Name", "Grade", "Room", "FAMILY_ID"], [f"[{EMMA}]", "3", "B", f"[{SMITHS}]"], [], ["Total", "3"]]
    res = _restore(out, key, out.name)
    assert res.output_path.read_bytes() == data


# ---------------------------------------------------------------------------
# 27. A formula name cell: the key records only rewrites that happened
# ---------------------------------------------------------------------------

def _formula_roster() -> bytes:
    """Full Name is a formula with a cached value, as Excel saves it."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Full Name", "First", "Last"])
    ws.append(['=B2&" "&C2', "Emma", "Smith"])
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode()
                xml, n = re.subn(r'<c r="A2"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>',
                                 r'<c r="A2"\1 t="str"><f>\2</f><v>Emma Smith</v></c>', xml)
                assert n == 1, xml
                data = xml.encode()
            zout.writestr(info, data)
    return out.getvalue()


def test_a_formula_name_cell_is_reported_and_left_out_of_the_key():
    from app.extractors import extract
    from app.config import UPLOADS_DIR
    probe = UPLOADS_DIR / "probe.xlsx"
    probe.write_bytes(_formula_roster())
    assert extract(probe).payload["tables"][0][1][0] == "Emma Smith", "the cached value is what FG sees"
    probe.unlink()
    result = {"committed": True, "import_runs": ["imp_formula00001"], "summary": {},
              "sheets": [{"rows": [
                  {"index": 0, "persons": [_person("0:0:0", "Emma", "Smith", [{"col": 0, "part": "full"}], EMMA)],
                   "family": {"key": "0:0:family", "display_name": "Smith Family", "cell": None,
                              "action": "new", "community_id": SMITHS}}]}]}
    sess = _session("formula.xlsx", _formula_roster(), result, flagged=("Emma", "Smith"))
    out, key = _scrub(sess)
    assert _grid(out)[1][0] == '=B2&" "&C2', "the formula is left alone"
    assert key["identity_cells"] == [], "no rewrite happened, so none is recorded"
    steps = [s["step"] for s in sess.scrub_steps if s["status"] == "err"]
    assert any("1 roster cell" in s for s in steps), sess.scrub_steps


# ---------------------------------------------------------------------------
# 28. A one-column roster has a header row
# ---------------------------------------------------------------------------

def test_a_single_column_roster_is_found():
    from app import community
    for header in ("Name", "Student Name"):
        sheets = community.build_sheets([[[header], ["Jane Doe"], ["John Roe"]]])
        assert len(sheets) == 1 and sheets[0]["header_row"] == 0 and sheets[0]["headers"] == [header]
    # A title line above the real one-column header is not the header.
    sheets = community.build_sheets([[["Student Roster"], ["Name"], ["Jane Doe"]]])
    assert sheets[0]["header_row"] == 1
    # A one-cell title above a wide header never wins over it.
    sheets = community.build_sheets([[["Student Directory"], ["First", "Last"], ["Jane", "Doe"]]])
    assert sheets[0]["header_row"] == 1
    # A single cell with no header word is still not a header.
    assert community.build_sheets([[["Jane Doe"], ["John Roe"]]]) == []


# ---------------------------------------------------------------------------
# 32. Family Graph's sheet / column / cell limits are checked up front
# ---------------------------------------------------------------------------

def test_family_graph_limits_are_mirrored_with_a_clear_message():
    from app import community
    small = [["First", "Last"], ["Jane", "Doe"]]
    with pytest.raises(community.CommunityError, match="26 roster sheets.*up to 25"):
        community.build_sheets([small] * 26)
    # Tables with no header words (a totals tab) do not count against it.
    numbers = [["1", "2"], ["3", "4"]]
    assert len(community.build_sheets([small] * 25 + [numbers] * 3)) == 25
    wide = [["Name"] + [f"c{i}" for i in range(300)], ["Jane"] + [""] * 300]
    with pytest.raises(community.CommunityError, match="301 columns.*up to 300"):
        community.build_sheets([wide])
    long_cell = [["Name", "Notes"], ["Jane", "x" * 4001]]
    with pytest.raises(community.CommunityError, match="4,000 characters"):
        community.build_sheets([long_cell])
